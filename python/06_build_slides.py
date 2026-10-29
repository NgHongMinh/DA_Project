"""Sinh bản trình chiếu từ số liệu đã đối chiếu và nội dung dành cho người đọc."""
from __future__ import annotations
import datetime as dt
import json
import re
import zipfile
from pathlib import Path
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from config import OUTPUT
from phase6_contract import CLAIMS
from presentation_assets import build_assets
from presentation_copy import pages
from publication_checks import assert_slides_clean

VALUE = {c.claim_id: c.value for c in CLAIMS}
SLIDES_DIR = OUTPUT / "slides"
DECK_PATH = SLIDES_DIR / "ASEAN_Renewables.pptx"
ASSETS = SLIDES_DIR / "assets"
W, H = 13.333, 7.5
INK, MUTED, ACCENT = "374151", "6B7280", "D55E00"
SOURCE = "Nguồn: Our World in Data, tổng hợp từ Ember và BP · 10 nước ASEAN · 2000–2021"


def vn(value, decimals=0):
    whole, _, frac = f"{abs(value):.{decimals}f}".partition(".")
    grouped = f"{int(whole):,}".replace(",", ".")
    return ("-" if value < 0 else "") + grouped + ("," + frac if decimals else "")


def pct(value, decimals=2):
    return vn(value * 100, decimals) + "%"


def wrap_text(body, width, size, bold=False):
    font = ImageFont.truetype(str(Path("C:/Windows/Fonts") / ("segoeuib.ttf" if bold else "segoeui.ttf")), round(size * 4))
    lines = []
    for paragraph in body.split("\n"):
        line = ""
        for word in paragraph.split():
            candidate = (line + " " + word).strip()
            if line and font.getlength(candidate) > (width * 72 - 8) * 4:
                lines.append(line)
                line = word
            else:
                line = candidate
        lines.append(line)
    return lines


def text(slide, x, y, w, h, body, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT):
    for half_points in range(int(size * 2), 18, -1):
        actual = half_points / 2
        lines = wrap_text(body, w, actual, bold)
        if len(lines) * actual * 1.24 + 5 <= h * 72:
            break
    else:
        raise ValueError(f"Khung quá nhỏ: {body[:80]}")
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = False
    frame.margin_left = frame.margin_right = Pt(4)
    frame.margin_top = frame.margin_bottom = Pt(2)
    for i, line in enumerate(lines):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.alignment = align
        p.line_spacing = Pt(actual * 1.24)
        p.space_before = p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        run.font.name = "Segoe UI"
        run.font.size = Pt(actual)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)
    return box


def rect(slide, x, y, w, h, color="E5E7EB"):
    shape = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    shape.shadow.inherit = False


def picture(slide, name, x, y, w, h):
    path = ASSETS / name
    with Image.open(path) as img:
        ratio = img.width / img.height
    width, height = min(w, h * ratio), min(h, w / ratio)
    shape = slide.shapes.add_picture(str(path), Inches(x + (w-width)/2), Inches(y+(h-height)/2), width=Inches(width))
    props = shape._element.xpath(".//p:cNvPr")[0]
    props.set("name", "Biểu đồ")
    props.set("descr", "Biểu đồ số liệu điện tái tạo theo phạm vi và chỉ tiêu ghi trong hình.")


def bullets(slide, lines, x=.72, y=2.1, w=11.9, row=.65, size=19):
    for i, line in enumerate(lines):
        text(slide, x, y+i*row, w, row-.06, "• " + line, size)


def stats(slide, items):
    width = 12.0 / len(items)
    for i, (number, caption, color) in enumerate(items):
        x = .67 + i * width
        text(slide, x, 1.98, width-.18, .62, number, 33, True, color, PP_ALIGN.CENTER)
        text(slide, x, 2.66, width-.18, .72, caption, 14, color=MUTED, align=PP_ALIGN.CENTER)


def build_slides(prs):
    for page, spec in enumerate(pages(VALUE, vn, pct), 1):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        kind = spec.get("kind", "bullets")
        if kind == "opening":
            text(s, .62, .72, 12, .35, "ĐỒ ÁN PHÂN TÍCH DỮ LIỆU", 13, True, ACCENT)
            text(s, .62, 1.38, 12, 1.48, spec["title"], 37, True)
            rect(s, .62, 3.1, 3.2, .035, ACCENT)
            text(s, .62, 3.65, 7.35, 1.25, spec["body"], 22)
            text(s, 8.3, 3.62, 4.3, 1.12, pct(VALUE["F4_VN_SHARE_MODERN"]), 61, True, ACCENT, PP_ALIGN.CENTER)
            text(s, 8.3, 4.87, 4.3, .9, "Tỷ trọng Việt Nam trong sản lượng điện gió và mặt trời của khu vực năm 2021", 15, color=MUTED, align=PP_ALIGN.CENTER)
            text(s, .62, 5.98, 12, .83, spec["glossary"], 13, color=MUTED)
        elif kind == "closing":
            text(s, .72, 2.45, 11.9, 1.1, "Xin cảm ơn", 54, True)
            rect(s, .72, 3.89, 3.2, .035, ACCENT)
            text(s, .72, 4.3, 11.9, .65, "Trao đổi và câu hỏi", 27, color=ACCENT)
            text(s, .72, 5.48, 11.9, .82, spec["body"], 17, color=MUTED)
        else:
            text(s, .62, .35, 12.09, .3, spec["category"].upper(), 12, True, ACCENT)
            text(s, .62, .78, 12.09, .94, spec["title"], 30, True)
            rect(s, .62, 1.79, 12.09, .018)
            if kind == "finding":
                picture(s, spec["image"], .68, 2.0, 8.28, 4.34)
                for i, (number, caption, color) in enumerate(spec["stats"]):
                    y = 2.04 + i * (1.48 if len(spec["stats"]) == 3 else 1.95)
                    text(s, 9.18, y, 3.45, .52, number, 31, True, color, PP_ALIGN.CENTER)
                    text(s, 9.18, y+.57, 3.45, .82, caption, 14, color=MUTED, align=PP_ALIGN.CENTER)
                text(s, .7, 6.43, 11.95, .51, spec["implication"], 15, True)
            elif kind == "quality":
                stats(s, spec["stats"])
                if spec.get("definition"):
                    text(s, .72, 3.58, 11.9, .5, spec["definition"], 18, True)
                bullets(s, spec["lines"], y=4.15, row=.62, size=17)
            elif kind == "limits":
                bullets(s, spec["lines"], w=5.3, row=.59, size=17)
                picture(s, "capacity_coverage.png", 6.22, 2.14, 6.42, 3.88)
                text(s, 6.22, 6.05, 6.42, .72, "Hệ số sử dụng công suất xấp xỉ có 104 quan sát nước–năm–công nghệ.", 15, color=MUTED)
            elif kind == "statistics":
                widths = [2.15, .65, .94, .84, 1.0]
                for r, row in enumerate(spec["table"]):
                    x = .68
                    for width, cell in zip(widths, row):
                        text(s, x, 2.12+r*.66, width, .61, cell, 13, r == 0, MUTED if r == 0 else INK)
                        x += width
                    rect(s, .68, 2.75+r*.66, sum(widths), .008)
                text(s, .68, 4.96, 5.65, .65, spec["percentiles"], 16)
                text(s, .68, 5.78, 5.65, .96, spec["implication"], 16, color=MUTED)
                picture(s, "solar_distribution.png", 6.64, 2.05, 6, 4.68)
            elif kind == "cards":
                for i, (title, body, color) in enumerate(spec["cards"]):
                    x = .72+i*4.07
                    rect(s, x, 2.34, 3.79, 2.47, "F3F4F6")
                    rect(s, x, 2.34, 3.79, .055, color)
                    text(s, x+.15, 2.6, 3.49, .8, title, 23, True, color)
                    text(s, x+.15, 3.56, 3.49, 1, body, 18)
                text(s, .75, 5.25, 11.82, 1.05, spec["implication"], 21)
            elif kind == "columns":
                for i, (title, lines) in enumerate(spec["columns"]):
                    x = .75+i*6.17
                    if title:
                        text(s, x, 2.1, 5.67, .75, title, 18, True, "0072B2" if i == 0 else ACCENT)
                    bullets(s, lines, x=x, y=3.05 if title else 2.15, w=5.67,
                            row=3.4/len(lines) if title else .76, size=18)
            elif kind == "needs":
                for i, (need, purpose) in enumerate(spec["rows"]):
                    y = 2.18+i*.95
                    text(s, .75, y, 5.66, .72, need, 20, True, "0072B2")
                    text(s, 6.87, y, 5.71, .72, purpose, 19)
                    rect(s, .75, y+.8, 11.83, .01)
            else:
                bullets(s, spec["lines"], y=2.06, row=.84 if page == 14 else .55, size=20 if page == 14 else 18)
                if spec.get("implication"):
                    rect(s, .72, 5.55, 11.9, .015)
                    text(s, .72, 5.78, 11.9, .89, spec["implication"], 17)
        s.notes_slide.notes_text_frame.text = spec["notes"].strip()
        text(s, .62, 7.05, 11.45, .24, SOURCE, 9.5, color=MUTED)
        text(s, 12.15, 7.05, .5, .24, str(page), 9.5, color=MUTED, align=PP_ALIGN.RIGHT)


# Nhãn năm, số trang, ngưỡng và các số đếm chất lượng dữ liệu đã đối chiếu.
ALLOWED_LITERALS = {str(i) for i in range(18)} | {"2000", "2010", "2021", "2022", "25", "75", "0,05", "3,125", "106", "247", "89.075", "864", "104"}


def assert_no_stray_numbers(prs):
    rendered = set(ALLOWED_LITERALS)
    derived = [VALUE["F1_ASEAN_2021"]/VALUE["F1_ASEAN_2000"],
               VALUE["F2_HYDRO_2021"]/VALUE["F1_ASEAN_2021"],
               VALUE["F2_MODERN_2021"]/VALUE["F1_ASEAN_2021"]]
    for value in [*VALUE.values(), *derived]:
        for decimals in range(5):
            rendered.update((vn(value, decimals), vn(value*100, decimals)))
    errors = []
    for i, slide in enumerate(prs.slides, 1):
        content = "\n".join(sh.text for sh in slide.shapes if sh.has_text_frame)
        content += "\n" + slide.notes_slide.notes_text_frame.text
        for token in re.findall(r"\d[\d.,]*", content):
            if token.rstrip(".,") not in rendered:
                errors.append(f"Trang {i}: {token}")
    return errors


def assert_layout(prs):
    for i, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if sh.left < 0 or sh.top < 0 or sh.left+sh.width > prs.slide_width+2 or sh.top+sh.height > prs.slide_height+2:
                raise AssertionError(f"Trang {i}: khung vượt mép trang")


def repack_deterministically(path):
    with zipfile.ZipFile(path) as z:
        parts = {name: z.read(name) for name in sorted(z.namelist())}
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, payload in parts.items():
            info = zipfile.ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            z.writestr(info, payload)


def main():
    build_assets(ASSETS, VALUE)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    build_slides(prs)
    assert len(prs.slides) == 17
    assert_layout(prs)
    assert_slides_clean(prs, json.loads((ASSETS / "figure_text.json").read_text(encoding="utf-8")))
    errors = assert_no_stray_numbers(prs)
    if errors:
        raise AssertionError("\n".join(errors))
    core = prs.core_properties
    core.title = "Điện tái tạo tại Đông Nam Á: quy mô, cơ cấu và vị trí của Việt Nam"
    core.author = core.last_modified_by = "Nguyễn Hồng Minh"
    core.comments = core.subject = core.keywords = ""
    core.created = core.modified = dt.datetime(2026, 1, 1)
    core.revision = 1
    prs.save(DECK_PATH)
    repack_deterministically(DECK_PATH)
    print("Đã tạo 17 trang; số liệu và vị trí hình đều được kiểm tra.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
