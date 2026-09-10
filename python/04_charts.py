"""Phase 4: build, validate, render, and atomically publish 13 charts.

The Phase 3 fact surface remains the authoritative analytical source.  This
module only reshapes verified facts for presentation; the sole new analytical
metric is the capacity-factor proxy documented for C13.
"""

from __future__ import annotations

import ast
import hashlib
import json
import math
import os
import shutil
import sqlite3
import sys
import tempfile
import textwrap
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from matplotlib.ticker import FuncFormatter, PercentFormatter
import numpy as np
import pandas as pd
import seaborn as sns
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageStat

from config import ARCHIVE, CHARTS, CLEAN, DB_PATH, OUTPUT, STATS
from phase2_contract import Phase2ContractError, validate_phase2_gate
from phase3_contract import PHASE3_WORKBOOKS
from phase4_contract import (
    CAPACITY_FACTOR_COUNTRIES,
    CAPACITY_FACTOR_TOLERANCE,
    CHART_BY_ID,
    CHART_DATA_SHEETS,
    CHART_FILES,
    CHART_SPECS,
    COUNTRY_COLORS,
    DIMENSIONS_PX,
    DPI,
    EXPECTED_C01_ENDPOINTS,
    EXPECTED_C02_2021_TWH,
    EXPECTED_C04_RANK_SHIFTS,
    EXPECTED_C08_OPPOSITE_DIRECTION_YEARS,
    EXPECTED_C09_2021_GLOBAL_RANKS,
    EXPECTED_C09_2021_TWH,
    EXPECTED_C10_OBSERVED_N,
    EXPECTED_C10_ZERO_N,
    EXPECTED_C12_2021_TOP3_SHARE,
    EXPECTED_C12_VIETNAM_POSITIVE_GROWTH_SHARE,
    EXPECTED_CAPACITY_FACTOR_ROWS,
    EXPECTED_PHASE2_FACT_ENERGY_SHA256,
    EXPECTED_PHASE2_FACT_ROWS,
    EXPECTED_PHASE2_FINGERPRINT,
    EXPECTED_PHASE2_BLOCKING_FAILURES,
    EXPECTED_PHASE2_STATUS_COUNTS,
    EXPECTED_PHASE2_VALIDATION_ROWS,
    EXPECTED_PHASE3_CROSSCHECK_ROWS,
    EXPECTED_PHASE3_FACT_ROWS,
    EXPECTED_PHASE3_VALIDATION_ROWS,
    EXPECTED_QUESTION_ROWS,
    EXPECTED_VALIDATION_IDS,
    FIGSIZE_INCHES,
    FONT_FAMILY,
    MANIFEST_COLUMNS,
    METHOD_VERSION,
    NUMERIC_TOLERANCE,
    QUESTION_IDS_BY_CHART,
    REQUIRED_METADATA_KEYS,
    REQUIRED_METHOD_NOTE_KEYS,
    SHEET_KEY_COLUMNS,
    SHEET_REQUIRED_COLUMNS,
    TECHNOLOGY_COLORS,
    VALIDATION_COLUMNS,
    VALIDATION_GATE_BY_ID,
    YEAR_END,
    YEAR_START,
    YEARS,
    manifest_rows,
)


FACT_KEY_COLUMNS = [
    "subject_type",
    "subject_key",
    "period_type",
    "period_key",
    "series_key",
    "metric_name",
    "comparison_scope",
    "filter_variant",
    "method_version",
]
ASEAN = [
    "Brunei",
    "Cambodia",
    "Indonesia",
    "Laos",
    "Malaysia",
    "Myanmar",
    "Philippines",
    "Singapore",
    "Thailand",
    "Vietnam",
]
TECHNOLOGIES = ["HYDRO", "OTHER_RENEWABLES", "WIND", "SOLAR"]
POSITION_SERIES = ["RENEWABLE_TOTAL", "MODERN_WIND_SOLAR"]
SERIES_LABELS = {
    "HYDRO": "Thủy điện",
    "OTHER_RENEWABLES": "NLTT khác",
    "WIND": "Gió",
    "SOLAR": "Mặt trời",
    "RENEWABLE_TOTAL": "Tổng NLTT",
    "MODERN_WIND_SOLAR": "Gió + mặt trời",
}
ENTITY_LABELS = {
    "Brunei": "Brunei",
    "Cambodia": "Campuchia",
    "Indonesia": "Indonesia",
    "Laos": "Lào",
    "Malaysia": "Malaysia",
    "Myanmar": "Myanmar",
    "Philippines": "Philippines",
    "Singapore": "Singapore",
    "Thailand": "Thái Lan",
    "Vietnam": "Việt Nam",
    "ASEAN_10": "ASEAN-10",
    "WORLD_REPORTED": "Thế giới",
}
PHASE3_PUBLIC_FILES = [
    STATS / "03_fact_metrics.csv",
    STATS / "03_metric_definitions.csv",
    STATS / "03_crosscheck_results.csv",
    STATS / "03_validation_results.csv",
    *(STATS / name for name in PHASE3_WORKBOOKS),
    DB_PATH,
]


class Phase4Error(RuntimeError):
    """Raised when a blocking Phase 4 contract fails."""


@dataclass
class Inputs:
    facts: pd.DataFrame
    definitions: pd.DataFrame
    phase2_validation: pd.DataFrame
    phase3_validation: pd.DataFrame
    crosscheck: pd.DataFrame
    fact_energy: pd.DataFrame
    fact_energy_sha256: str
    fingerprint: str
    workbook_fact_keys: set[str]
    workbook_fingerprints: set[str]
    immutable_hashes: dict[str, str]


class ValidationCollector:
    def __init__(self, fingerprint: str) -> None:
        self.fingerprint = fingerprint
        self.rows: list[dict[str, Any]] = []

    def add(
        self,
        validation_id: str,
        validation_group: str,
        chart_id: str,
        condition: bool,
        expected: Any,
        actual: Any,
        details: str,
    ) -> None:
        self.rows.append(
            {
                "validation_id": validation_id,
                "validation_group": validation_group,
                "gate_id": VALIDATION_GATE_BY_ID[validation_id],
                "chart_id": chart_id,
                "question_id": (
                    ",".join(QUESTION_IDS_BY_CHART[chart_id])
                    if chart_id in QUESTION_IDS_BY_CHART
                    else "NOT_APPLICABLE"
                ),
                "severity": "BLOCKING",
                "status": "PASS" if bool(condition) else "FAIL",
                "expected_value": _stable_text(expected),
                "actual_value": _stable_text(actual),
                "failure_reason": "" if bool(condition) else validation_id,
                "details": details,
                "method_version": METHOD_VERSION,
                "phase2_input_fingerprint": self.fingerprint,
            }
        )

    def frame(self) -> pd.DataFrame:
        frame = pd.DataFrame(self.rows)
        missing = set(VALIDATION_COLUMNS) - set(frame.columns)
        for column in missing:
            frame[column] = ""
        return frame[list(VALIDATION_COLUMNS)]


def _stable_text(value: Any) -> str:
    def normalize(item: Any) -> Any:
        if isinstance(item, np.generic):
            return item.item()
        if isinstance(item, Path):
            return str(item)
        if isinstance(item, dict):
            return {str(key): normalize(entry) for key, entry in item.items()}
        if isinstance(item, (list, tuple)):
            return [normalize(entry) for entry in item]
        if isinstance(item, set):
            return sorted(normalize(entry) for entry in item)
        return item

    if isinstance(value, (dict, list, tuple, set)):
        normalized = normalize(value)
        return json.dumps(normalized, ensure_ascii=False, separators=(",", ":"))
    if isinstance(value, float):
        if math.isnan(value):
            return "NaN"
        return format(value, ".12g")
    return str(value)


def file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def file_hashes(paths: Iterable[Path]) -> dict[str, str]:
    return {str(path.resolve()): file_hash(path) for path in paths}


def add_fact_key(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.copy()
    result["fact_key"] = result[FACT_KEY_COLUMNS].apply(
        lambda row: json.dumps(row.tolist(), ensure_ascii=False, separators=(",", ":")),
        axis=1,
    )
    return result


def select_facts(facts: pd.DataFrame, **conditions: Any) -> pd.DataFrame:
    selected = facts
    for column, accepted in conditions.items():
        values = accepted if isinstance(accepted, (list, tuple, set, np.ndarray)) else [accepted]
        selected = selected[selected[column].isin(values)]
    return add_fact_key(selected.copy()).reset_index(drop=True)


def _period_to_year(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.copy()
    if "period_start_year" in result:
        result["year"] = pd.to_numeric(result["period_start_year"], errors="coerce").astype("Int64")
    return result


def _collect_workbook_fact_keys() -> tuple[set[str], bool, set[str]]:
    keys: set[str] = set()
    fingerprints: set[str] = set()
    contracts_ok = True
    for file_name, expected_sheets in PHASE3_WORKBOOKS.items():
        path = STATS / file_name
        if not path.exists():
            return keys, False, fingerprints
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            if workbook.sheetnames != expected_sheets:
                contracts_ok = False
            for worksheet in workbook.worksheets:
                rows = worksheet.iter_rows(values_only=True)
                try:
                    header = [str(value) if value is not None else "" for value in next(rows)]
                except StopIteration:
                    continue
                key_indexes = [
                    index
                    for index, column in enumerate(header)
                    if column == "fact_key" or column.endswith("_fact_key")
                ]
                fingerprint_indexes = [
                    index
                    for index, column in enumerate(header)
                    if column == "phase2_input_fingerprint"
                ]
                for row in rows:
                    for index in key_indexes:
                        value = row[index]
                        if isinstance(value, str) and value.startswith("["):
                            keys.add(value)
                    for index in fingerprint_indexes:
                        value = row[index]
                        if isinstance(value, str) and value:
                            fingerprints.add(value)
        finally:
            workbook.close()
    return keys, contracts_ok, fingerprints


def load_inputs() -> tuple[Inputs, dict[str, Any]]:
    required = [
        STATS / "02_validation_results.csv",
        STATS / "03_fact_metrics.csv",
        STATS / "03_metric_definitions.csv",
        STATS / "03_crosscheck_results.csv",
        STATS / "03_validation_results.csv",
        CLEAN / "fact_energy.csv",
        *PHASE3_PUBLIC_FILES,
        *sorted(ARCHIVE.glob("*.csv")),
    ]
    missing = sorted({str(path) for path in required if not path.exists()})
    if missing:
        raise Phase4Error(f"Missing required Phase 2/3 inputs: {missing}")

    fact_energy_path = CLEAN / "fact_energy.csv"
    immutable_paths = [*sorted(ARCHIVE.glob("*.csv")), fact_energy_path, *PHASE3_PUBLIC_FILES]
    immutable_hashes = file_hashes(immutable_paths)
    facts = pd.read_csv(STATS / "03_fact_metrics.csv", low_memory=False)
    facts = add_fact_key(facts)
    definitions = pd.read_csv(STATS / "03_metric_definitions.csv", low_memory=False)
    phase2_validation = pd.read_csv(STATS / "02_validation_results.csv", low_memory=False)
    phase3_validation = pd.read_csv(STATS / "03_validation_results.csv", low_memory=False)
    crosscheck = pd.read_csv(STATS / "03_crosscheck_results.csv", low_memory=False)
    fact_energy = pd.read_csv(fact_energy_path, low_memory=False)
    fact_energy_sha256 = immutable_hashes[str(fact_energy_path.resolve())]
    phase2_gate_error = ""
    phase2_db_fingerprint = ""
    try:
        connection = sqlite3.connect(DB_PATH.resolve().as_uri() + "?mode=ro", uri=True)
        try:
            phase2_db_fingerprint = validate_phase2_gate(connection)
        finally:
            connection.close()
    except (Phase2ContractError, sqlite3.Error, OSError) as exc:
        phase2_gate_error = str(exc)
    fingerprints = set(facts["phase2_input_fingerprint"].dropna().astype(str))
    fingerprint = next(iter(fingerprints)) if len(fingerprints) == 1 else ""
    workbook_keys, workbook_contracts_ok, workbook_fingerprints = _collect_workbook_fact_keys()

    question_counts = facts.groupby("question_id").size().sort_index().to_dict()
    phase2_status = phase2_validation.groupby("status").size().to_dict()
    phase2_blocking = int(
        (
            pd.to_numeric(phase2_validation["is_blocking"], errors="coerce").fillna(0).astype(int).eq(1)
            & phase2_validation["status"].ne("PASS")
        ).sum()
    )
    input_checks = {
        "phase2_rows": len(phase2_validation),
        "phase2_status": phase2_status,
        "phase2_blocking": phase2_blocking,
        "phase2_fact_rows": len(fact_energy),
        "phase2_fact_energy_sha256": fact_energy_sha256,
        "phase2_db_fingerprint": phase2_db_fingerprint,
        "phase2_gate_error": phase2_gate_error,
        "phase3_fact_rows": len(facts),
        "phase3_question_rows": question_counts,
        "phase3_cross_rows": len(crosscheck),
        "phase3_cross_pass": int(crosscheck["status"].eq("PASS").sum()),
        "phase3_validation_rows": len(phase3_validation),
        "phase3_validation_pass": int(phase3_validation["status"].eq("PASS").sum()),
        "fingerprints": fingerprints,
        "workbook_contracts_ok": workbook_contracts_ok,
        "archive_count": len(list(ARCHIVE.glob("*.csv"))),
    }
    inputs = Inputs(
        facts=facts,
        definitions=definitions,
        phase2_validation=phase2_validation,
        phase3_validation=phase3_validation,
        crosscheck=crosscheck,
        fact_energy=fact_energy,
        fact_energy_sha256=fact_energy_sha256,
        fingerprint=fingerprint,
        workbook_fact_keys=workbook_keys,
        workbook_fingerprints=workbook_fingerprints,
        immutable_hashes=immutable_hashes,
    )
    return inputs, input_checks


def _chart_frame(frame: pd.DataFrame, chart_id: str) -> pd.DataFrame:
    result = _period_to_year(frame)
    if "method_version" in result.columns:
        result = result.rename(columns={"method_version": "source_method_version"})
    result.insert(0, "chart_id", chart_id)
    result["method_version"] = METHOD_VERSION
    return result


def _merge_rank_data(facts: pd.DataFrame, fingerprint: str) -> pd.DataFrame:
    base = {
        "subject_type": "COUNTRY",
        "period_type": "YEAR",
        "period_key": "Y2021",
        "comparison_scope": "ASEAN_10",
        "filter_variant": "NOT_APPLICABLE",
    }
    total = select_facts(
        facts,
        **base,
        series_key="RENEWABLE_TOTAL",
        metric_name="asean_rank",
    )[["subject_key", "metric_value", "unit", "value_status", "exclusion_reason", "fact_key"]].rename(
        columns={
            "metric_value": "total_rank",
            "unit": "total_rank_unit",
            "value_status": "total_rank_status",
            "exclusion_reason": "total_rank_exclusion_reason",
            "fact_key": "total_rank_fact_key",
        }
    )
    modern = select_facts(
        facts,
        **base,
        series_key="MODERN_WIND_SOLAR",
        metric_name="asean_rank",
    )[["subject_key", "metric_value", "unit", "value_status", "exclusion_reason", "fact_key"]].rename(
        columns={
            "metric_value": "modern_rank",
            "unit": "modern_rank_unit",
            "value_status": "modern_rank_status",
            "exclusion_reason": "modern_rank_exclusion_reason",
            "fact_key": "modern_rank_fact_key",
        }
    )
    gain = select_facts(
        facts,
        **base,
        series_key="MODERN_WIND_SOLAR",
        metric_name="modern_rank_gain",
    )[["subject_key", "metric_value", "unit", "value_status", "exclusion_reason", "fact_key"]].rename(
        columns={
            "metric_value": "modern_rank_gain",
            "unit": "modern_rank_gain_unit",
            "value_status": "modern_rank_gain_status",
            "exclusion_reason": "modern_rank_gain_exclusion_reason",
            "fact_key": "modern_rank_gain_fact_key",
        }
    )
    result = total.merge(modern, on="subject_key", validate="one_to_one").merge(
        gain, on="subject_key", validate="one_to_one"
    )
    result = result.rename(columns={"subject_key": "entity"})
    result.insert(0, "chart_id", "C04")
    result["year"] = 2021
    result["method_version"] = METHOD_VERSION
    result["phase2_input_fingerprint"] = fingerprint
    return result


def _merge_q5_data(facts: pd.DataFrame, fingerprint: str) -> pd.DataFrame:
    common = {
        "subject_type": "COUNTRY",
        "subject_key": "Vietnam",
        "period_type": "YEAR",
        "series_key": POSITION_SERIES,
        "comparison_scope": "ASEAN_10",
        "filter_variant": "NOT_APPLICABLE",
    }
    parts: list[pd.DataFrame] = []
    names = {
        "generation_twh": "generation_twh",
        "share_electricity_pct": "share_electricity_pct",
        "direction_relation": "direction_relation",
    }
    for metric_name, prefix in names.items():
        part = select_facts(facts, **common, metric_name=metric_name)
        columns = ["period_start_year", "series_key", "metric_value", "metric_text_value", "unit", "value_status", "exclusion_reason", "fact_key"]
        part = part[columns].rename(
            columns={
                "period_start_year": "year",
                "metric_value": prefix,
                "metric_text_value": f"{prefix}_text",
                "unit": f"{prefix}_unit",
                "value_status": f"{prefix}_status",
                "exclusion_reason": f"{prefix}_exclusion_reason",
                "fact_key": f"{prefix}_fact_key",
            }
        )
        parts.append(part)
    result = parts[0]
    for part in parts[1:]:
        result = result.merge(part, on=["year", "series_key"], validate="one_to_one")
    result.insert(0, "chart_id", "C08")
    result["year"] = result["year"].astype(int)
    result["method_version"] = METHOD_VERSION
    result["phase2_input_fingerprint"] = fingerprint
    return result


def _build_vietnam_position(facts: pd.DataFrame, fingerprint: str) -> pd.DataFrame:
    generation = select_facts(
        facts,
        subject_type="COUNTRY",
        subject_key="Vietnam",
        period_type="YEAR",
        series_key=POSITION_SERIES,
        metric_name="generation_twh",
        comparison_scope="ASEAN_10",
        filter_variant="NOT_APPLICABLE",
    )
    generation = _period_to_year(generation)
    generation = generation.assign(record_type="TIME_SERIES", display_subject="Vietnam")

    asean9 = select_facts(
        facts,
        subject_type="COUNTRY",
        subject_key="Vietnam",
        period_type="YEAR",
        series_key=POSITION_SERIES,
        metric_name="asean9_sum_twh",
        comparison_scope="ASEAN_9",
        filter_variant="NOT_APPLICABLE",
    )
    asean9 = _period_to_year(asean9)
    asean9 = asean9.assign(record_type="TIME_SERIES", display_subject="ASEAN_9_TOTAL")

    ranks = select_facts(
        facts,
        subject_type="COUNTRY",
        subject_key="Vietnam",
        period_type="YEAR",
        period_key="Y2021",
        series_key=POSITION_SERIES,
        metric_name="global_rank",
        comparison_scope="GLOBAL_CURRENT_COUNTRIES",
        filter_variant="NOT_APPLICABLE",
    )
    ranks = _period_to_year(ranks)
    ranks = ranks.assign(record_type="GLOBAL_RANK", display_subject="Vietnam")
    keep = [
        "record_type",
        "display_subject",
        "year",
        "series_key",
        "metric_name",
        "metric_value",
        "unit",
        "value_status",
        "exclusion_reason",
        "ranked_n",
        "fact_key",
        "method_version",
        "phase2_input_fingerprint",
    ]
    result = pd.concat([generation[keep], asean9[keep], ranks[keep]], ignore_index=True)
    result.insert(0, "chart_id", "C09")
    result = result.rename(columns={"method_version": "source_method_version"})
    result["method_version"] = METHOD_VERSION
    result["phase2_input_fingerprint"] = fingerprint
    return result


def _build_global_distribution(facts: pd.DataFrame, fingerprint: str) -> pd.DataFrame:
    values = select_facts(
        facts,
        subject_type="COUNTRY",
        period_type="YEAR",
        period_key="Y2021",
        series_key=TECHNOLOGIES,
        metric_name="generation_twh",
        comparison_scope="GLOBAL_CURRENT_COUNTRIES",
        filter_variant="NOT_APPLICABLE",
    )
    values = _period_to_year(values)
    numeric = pd.to_numeric(values["metric_value"], errors="coerce")
    threshold = pd.to_numeric(values["threshold_value"], errors="coerce").fillna(0.05)
    values["display_class"] = np.select(
        [numeric.isna(), numeric.eq(0), numeric.gt(0) & numeric.lt(threshold), numeric.ge(threshold)],
        ["Missing", "Zero", "Trace", "Material"],
        default="Missing",
    )
    values["record_type"] = "COUNTRY_VALUE"

    summary = select_facts(
        facts,
        subject_type="STATISTICAL_SCOPE",
        period_type="YEAR",
        period_key="Y2021",
        series_key=TECHNOLOGIES,
        metric_name="n",
        comparison_scope="GLOBAL_CURRENT_COUNTRIES",
        filter_variant="ALL_OBSERVED",
    )
    summary = _period_to_year(summary)
    summary["display_class"] = "SUMMARY"
    summary["record_type"] = "SUMMARY"
    common = sorted(set(values.columns) | set(summary.columns))
    result = pd.concat([values.reindex(columns=common), summary.reindex(columns=common)], ignore_index=True)
    result.insert(0, "chart_id", "C10")
    result = result.rename(columns={"method_version": "source_method_version"})
    result["method_version"] = METHOD_VERSION
    result["phase2_input_fingerprint"] = fingerprint
    return result


def _build_entry_lag(facts: pd.DataFrame, fingerprint: str) -> pd.DataFrame:
    common = {
        "subject_type": "COUNTRY",
        "period_type": "WINDOW",
        "series_key": TECHNOLOGIES,
        "comparison_scope": "ASEAN_10",
        "filter_variant": "NOT_APPLICABLE",
    }
    status = select_facts(facts, **common, metric_name="material_entry_status")
    year = select_facts(facts, **common, metric_name="observed_material_entry_year")
    status = status[["subject_key", "series_key", "metric_text_value", "unit", "value_status", "exclusion_reason", "fact_key"]].rename(
        columns={
            "subject_key": "entity",
            "metric_text_value": "material_entry_status_text",
            "unit": "material_entry_status_unit",
            "value_status": "material_entry_status_value_status",
            "exclusion_reason": "material_entry_status_exclusion_reason",
            "fact_key": "material_entry_status_fact_key",
        }
    )
    year = year[["subject_key", "series_key", "metric_value", "unit", "value_status", "exclusion_reason", "fact_key"]].rename(
        columns={
            "subject_key": "entity",
            "metric_value": "observed_material_entry_year",
            "unit": "observed_material_entry_year_unit",
            "value_status": "observed_material_entry_year_value_status",
            "exclusion_reason": "observed_material_entry_year_exclusion_reason",
            "fact_key": "observed_material_entry_year_fact_key",
        }
    )
    result = status.merge(year, on=["entity", "series_key"], validate="one_to_one")
    result.insert(0, "chart_id", "C11")
    result["threshold_twh"] = 0.05
    result["method_version"] = METHOD_VERSION
    result["phase2_input_fingerprint"] = fingerprint
    return result


def _build_growth_data(facts: pd.DataFrame, fingerprint: str) -> pd.DataFrame:
    countries = select_facts(
        facts,
        subject_type="COUNTRY",
        period_type="INTERVAL",
        period_key="I2010_2021",
        series_key="MODERN_WIND_SOLAR",
        metric_name="positive_growth_share",
        comparison_scope="ASEAN_10",
        filter_variant="NOT_APPLICABLE",
    )
    aggregate = select_facts(
        facts,
        subject_type="REGION",
        subject_key="ASEAN_10",
        period_type="INTERVAL",
        period_key="I2010_2021",
        series_key="MODERN_WIND_SOLAR",
        metric_name=["gross_positive_addition_twh", "growth_positive_n", "top3_positive_growth_share"],
        comparison_scope="ASEAN_10",
        filter_variant="NOT_APPLICABLE",
    )
    result = pd.concat([countries, aggregate], ignore_index=True)
    result.insert(0, "chart_id", "C12")
    result = result.rename(columns={"method_version": "source_method_version"})
    result["method_version"] = METHOD_VERSION
    result["phase2_input_fingerprint"] = fingerprint
    return result


def _build_capacity_factor(fact_energy: pd.DataFrame, fingerprint: str) -> pd.DataFrame:
    rows: list[pd.DataFrame] = []
    for technology in ["wind", "solar"]:
        generation_key = f"{technology}_generation_twh"
        capacity_key = f"{technology}_capacity_gw"
        generation = fact_energy[
            fact_energy["entity"].isin(ASEAN)
            & fact_energy["year"].between(YEAR_START, YEAR_END)
            & fact_energy["metric_key"].eq(generation_key)
        ][["entity", "year", "value", "value_class", "source_file"]].rename(
            columns={
                "value": "generation_twh",
                "value_class": "generation_value_class",
                "source_file": "generation_source_file",
            }
        )
        capacity = fact_energy[
            fact_energy["entity"].isin(ASEAN)
            & fact_energy["year"].between(YEAR_START, YEAR_END)
            & fact_energy["metric_key"].eq(capacity_key)
        ][["entity", "year", "value", "value_class", "source_file"]].rename(
            columns={
                "value": "capacity_gw",
                "value_class": "capacity_value_class",
                "source_file": "capacity_source_file",
            }
        )
        merged = generation.merge(capacity, on=["entity", "year"], validate="one_to_one")
        merged["generation_twh"] = pd.to_numeric(merged["generation_twh"], errors="coerce")
        merged["capacity_gw"] = pd.to_numeric(merged["capacity_gw"], errors="coerce")
        merged = merged[
            merged["generation_twh"].notna()
            & np.isfinite(merged["generation_twh"])
            & merged["capacity_gw"].gt(0)
            & np.isfinite(merged["capacity_gw"])
        ].copy()
        merged["technology"] = technology.upper()
        merged["capacity_factor_ratio"] = merged["generation_twh"] / (merged["capacity_gw"] * 8.76)
        merged["generation_source_key"] = merged.apply(
            lambda row: json.dumps([row["entity"], int(row["year"]), generation_key], ensure_ascii=False, separators=(",", ":")),
            axis=1,
        )
        merged["capacity_source_key"] = merged.apply(
            lambda row: json.dumps([row["entity"], int(row["year"]), capacity_key], ensure_ascii=False, separators=(",", ":")),
            axis=1,
        )
        rows.append(merged)
    result = pd.concat(rows, ignore_index=True).sort_values(["technology", "entity", "year"]).reset_index(drop=True)
    result.insert(0, "chart_id", "C13")
    result["value_status"] = "CALCULATED"
    result["exclusion_reason"] = "NONE"
    result["denominator_basis"] = "YEAR_END_INSTALLED_CAPACITY_GW_X_8_76"
    result["method_version"] = METHOD_VERSION
    result["phase2_input_fingerprint"] = fingerprint
    return result


def build_chart_data(inputs: Inputs) -> dict[str, pd.DataFrame]:
    facts = inputs.facts
    fingerprint = inputs.fingerprint
    global_context = pd.concat(
        [
            select_facts(
                facts,
                subject_key=["ASEAN_10", "WORLD_REPORTED"],
                period_type="YEAR",
                series_key="RENEWABLE_TOTAL",
                metric_name="generation_twh",
                filter_variant="NOT_APPLICABLE",
            ),
            select_facts(
                facts,
                subject_key="ASEAN_10",
                period_type="YEAR",
                series_key="RENEWABLE_TOTAL",
                metric_name="asean_share_world",
                comparison_scope="WORLD_REPORTED",
                filter_variant="NOT_APPLICABLE",
            ),
        ],
        ignore_index=True,
    )
    global_context = _chart_frame(global_context, "C01")

    asean_technology = _chart_frame(
        select_facts(
            facts,
            subject_key="ASEAN_10",
            period_type="YEAR",
            series_key=TECHNOLOGIES,
            metric_name="generation_twh",
            comparison_scope="ASEAN_10",
            filter_variant="NOT_APPLICABLE",
        ),
        "C02",
    )
    country_mix = _chart_frame(
        select_facts(
            facts,
            subject_type="COUNTRY",
            subject_key=ASEAN,
            period_type="YEAR",
            period_key="Y2021",
            series_key=TECHNOLOGIES,
            metric_name=["generation_twh", "technology_share"],
            comparison_scope="ASEAN_10",
            filter_variant="NOT_APPLICABLE",
        ),
        "C03",
    )
    country_mix["entity"] = country_mix["subject_key"]
    rank_shift = _merge_rank_data(facts, fingerprint)
    small_multiples = _chart_frame(
        select_facts(
            facts,
            subject_type="COUNTRY",
            subject_key=ASEAN,
            period_type="YEAR",
            series_key="RENEWABLE_TOTAL",
            metric_name="generation_twh",
            comparison_scope="ASEAN_10",
            filter_variant="NOT_APPLICABLE",
        ),
        "C05|C06",
    )
    small_multiples["entity"] = small_multiples["subject_key"]
    vietnam_technology = _chart_frame(
        select_facts(
            facts,
            subject_type="COUNTRY",
            subject_key="Vietnam",
            period_type="YEAR",
            series_key=TECHNOLOGIES,
            metric_name="generation_twh",
            comparison_scope="ASEAN_10",
            filter_variant="NOT_APPLICABLE",
        ),
        "C07",
    )
    vietnam_q5 = _merge_q5_data(facts, fingerprint)
    vietnam_position = _build_vietnam_position(facts, fingerprint)
    global_distribution = _build_global_distribution(facts, fingerprint)
    global_distribution["entity"] = global_distribution["subject_key"]
    entry_lag = _build_entry_lag(facts, fingerprint)
    concentration = _chart_frame(
        select_facts(
            facts,
            subject_type="REGION",
            subject_key="ASEAN_10",
            period_type="YEAR",
            series_key=POSITION_SERIES,
            metric_name=["top3_share", "hhi", "gini"],
            comparison_scope="ASEAN_10",
            filter_variant="NOT_APPLICABLE",
        ),
        "C12",
    )
    growth_contribution = _build_growth_data(facts, fingerprint)
    growth_contribution["entity"] = growth_contribution["subject_key"]
    capacity_factor = _build_capacity_factor(inputs.fact_energy, fingerprint)

    manifest = pd.DataFrame(manifest_rows(), columns=MANIFEST_COLUMNS)
    manifest["phase2_input_fingerprint"] = fingerprint
    archive_hashes = {
        Path(path).name: digest
        for path, digest in sorted(inputs.immutable_hashes.items())
        if Path(path).parent.resolve() == ARCHIVE.resolve()
    }
    metadata = pd.DataFrame(
        [
            ("method_version", METHOD_VERSION),
            ("phase2_input_fingerprint", fingerprint),
            ("year_start", YEAR_START),
            ("year_end", YEAR_END),
            ("chart_count", len(CHART_SPECS)),
            ("image_width_px", DIMENSIONS_PX[0]),
            ("image_height_px", DIMENSIONS_PX[1]),
            ("dpi", DPI),
            ("primary_fact_source", "output/stats/03_fact_metrics.csv"),
            ("definition_source", "output/stats/03_metric_definitions.csv"),
            ("capacity_factor_source", "output/clean/fact_energy.csv"),
            ("phase2_fact_energy_sha256", inputs.fact_energy_sha256),
            ("fact_key_columns", json.dumps(FACT_KEY_COLUMNS, ensure_ascii=False, separators=(",", ":"))),
            ("dataviz_skill_status", "UNAVAILABLE_INTERNAL_CONTRACT_USED"),
            ("archive_csv_sha256", json.dumps(archive_hashes, ensure_ascii=False, separators=(",", ":"))),
        ],
        columns=["metadata_key", "metadata_value"],
    )
    metadata["method_version"] = METHOD_VERSION
    metadata["phase2_input_fingerprint"] = fingerprint
    method_notes = pd.DataFrame(
        [
            ("NO_PHASE3_RECALCULATION", "C01-C12 reshape verified Phase 3 facts; no Phase 3 formula is redefined."),
            ("LOG_ZERO", "Zero values remain explicit counts and are never placed on a logarithmic axis."),
            ("CAPACITY_FACTOR", "generation_twh / (year_end_capacity_gw * 8.76); a coverage-limited proxy, not technical efficiency."),
            ("ENTRY_CENSORING", "Window-start and not-material-by-2021 statuses are symbols, never fabricated entry years."),
            ("RANK_SCOPE", "Global rank is shown only for 2021 because ranked_n changes over time."),
        ],
        columns=["note_key", "note"],
    )
    method_notes["method_version"] = METHOD_VERSION
    method_notes["phase2_input_fingerprint"] = fingerprint

    data = {
        "manifest": manifest,
        "global_context": global_context,
        "asean_technology": asean_technology,
        "country_mix_2021": country_mix,
        "rank_shift_2021": rank_shift,
        "small_multiples": small_multiples,
        "vietnam_technology": vietnam_technology,
        "vietnam_q5": vietnam_q5,
        "vietnam_position": vietnam_position,
        "global_distribution": global_distribution,
        "entry_lag": entry_lag,
        "concentration": concentration,
        "growth_contribution": growth_contribution,
        "capacity_factor": capacity_factor,
        "metadata": metadata,
        "method_notes": method_notes,
    }
    if list(data) != list(CHART_DATA_SHEETS):
        raise Phase4Error(f"Chart-data sheet order differs from contract: {list(data)}")
    return data


def _configure_plotting() -> None:
    # Seaborn's jittered point placement in C10 uses NumPy's generator.  Reset
    # it for every full render so decoded pixels are reproducible across runs.
    np.random.seed(20210906)
    sns.set_theme(style="whitegrid", font=FONT_FAMILY)
    plt.rcParams.update(
        {
            "font.family": FONT_FAMILY,
            "figure.facecolor": "white",
            "axes.facecolor": "white",
            "axes.edgecolor": "#CBD5E1",
            "axes.labelcolor": "#334155",
            "axes.titlecolor": "#0F172A",
            "text.color": "#0F172A",
            "xtick.color": "#475569",
            "ytick.color": "#475569",
            "grid.color": "#E2E8F0",
            "grid.linewidth": 0.7,
            "axes.titleweight": "bold",
            "axes.titlesize": 12,
            "axes.labelsize": 10,
            "xtick.labelsize": 9,
            "ytick.labelsize": 9,
            "legend.frameon": False,
        }
    )


def _new_figure(*args: Any, **kwargs: Any) -> tuple[plt.Figure, Any]:
    kwargs.setdefault("figsize", FIGSIZE_INCHES)
    kwargs.setdefault("dpi", DPI)
    return plt.subplots(*args, **kwargs)


def _decorate_figure(fig: plt.Figure, chart_id: str, footer: str | None = None) -> None:
    spec = CHART_BY_ID[chart_id]
    fig.suptitle(spec["title_vi"], x=0.06, y=0.965, ha="left", va="top", fontsize=19, fontweight="bold")
    fig.text(0.06, 0.91, spec["message_vi"], ha="left", va="top", fontsize=11.5, color="#475569")
    note = footer or (
        f"Nguồn: {spec['source_locator']}  |  Phạm vi: {spec['scope']}  |  "
        f"Giai đoạn: {spec['period']}  |  {spec['limitation_vi']}"
    )
    fig.text(
        0.06,
        0.025,
        textwrap.fill(note, width=185),
        ha="left",
        va="bottom",
        fontsize=8.2,
        color="#64748B",
    )


def _save_figure(fig: plt.Figure, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(
        path,
        dpi=DPI,
        facecolor="white",
        edgecolor="white",
        bbox_inches=None,
        metadata={"Software": METHOD_VERSION},
    )
    plt.close(fig)


def _style_axis(axis: plt.Axes, *, y_grid: bool = True) -> None:
    axis.spines[["top", "right"]].set_visible(False)
    if y_grid:
        axis.grid(True, axis="y", alpha=0.75)
    else:
        axis.grid(False, axis="y")
    axis.grid(False, axis="x")


def _set_year_axis(axis: plt.Axes) -> None:
    axis.set_xlim(YEAR_START, YEAR_END)
    axis.set_xticks([2000, 2005, 2010, 2015, 2021])
    axis.set_xlabel("")


def _vi_number(value: float, decimals: int = 1) -> str:
    rendered = f"{value:,.{decimals}f}"
    return rendered.replace(",", "X").replace(".", ",").replace("X", ".")


def _format_share_label(value: Any) -> str:
    """Format a ratio as a Vietnamese percentage without hiding positive values as zero."""
    if pd.isna(value):
        return ""
    numeric = float(value)
    if 0 < numeric < 0.0005:
        return "<0,1%"
    return f"{_vi_number(numeric * 100, 1)}%"


def _render_c01(data: pd.DataFrame, target: Path) -> None:
    fig, axes = _new_figure(2, 1, gridspec_kw={"height_ratios": [2.1, 1]}, sharex=True)
    generation = data[data["metric_name"].eq("generation_twh")]
    colors = {"ASEAN_10": "#D55E00", "WORLD_REPORTED": "#374151"}
    for subject in ["WORLD_REPORTED", "ASEAN_10"]:
        series = generation[generation["subject_key"].eq(subject)].sort_values("year")
        axes[0].plot(
            series["year"],
            series["metric_value"],
            color=colors[subject],
            linewidth=2.8,
            marker="o",
            markevery=[0, -1],
            label=ENTITY_LABELS[subject],
        )
        end = series.iloc[-1]
        axes[0].annotate(
            f"{ENTITY_LABELS[subject]}  {_vi_number(end['metric_value'], 1)} TWh",
            (end["year"], end["metric_value"]),
            xytext=(-5, 8 if subject == "WORLD_REPORTED" else -16),
            textcoords="offset points",
            ha="right",
            color=colors[subject],
            fontsize=9,
            fontweight="bold",
        )
    axes[0].set_yscale("log")
    axes[0].set_ylabel("Sản lượng (TWh, thang logarit)")
    axes[0].legend(loc="upper left", ncol=2)
    axes[0].yaxis.set_major_formatter(FuncFormatter(lambda value, _: _vi_number(value, 0)))
    _style_axis(axes[0])

    share = data[data["metric_name"].eq("asean_share_world")].sort_values("year")
    share_pct = share["metric_value"] * 100
    axes[1].plot(share["year"], share_pct, color="#D55E00", linewidth=2.8)
    axes[1].fill_between(share["year"], share_pct, color="#D55E00", alpha=0.12)
    axes[1].scatter([share.iloc[0]["year"], share.iloc[-1]["year"]], [share_pct.iloc[0], share_pct.iloc[-1]], color="#D55E00", zorder=3)
    axes[1].annotate("2,55%", (2000, share_pct.iloc[0]), xytext=(6, 7), textcoords="offset points", fontsize=9)
    axes[1].annotate("3,67%", (2021, share_pct.iloc[-1]), xytext=(-4, 7), textcoords="offset points", ha="right", fontsize=9, fontweight="bold")
    axes[1].set_ylabel("ASEAN / thế giới")
    axes[1].yaxis.set_major_formatter(FuncFormatter(lambda value, _: f"{_vi_number(value, 1)}%"))
    axes[1].set_ylim(0, max(5, share_pct.max() * 1.25))
    _set_year_axis(axes[1])
    _style_axis(axes[1])
    fig.subplots_adjust(left=0.09, right=0.96, top=0.84, bottom=0.15, hspace=0.12)
    _decorate_figure(fig, "C01")
    _save_figure(fig, target)


def _stacked_area(axis: plt.Axes, data: pd.DataFrame) -> None:
    pivot = data.pivot(index="year", columns="series_key", values="metric_value").sort_index()
    values = [pivot[series].to_numpy(dtype=float) for series in TECHNOLOGIES]
    axis.stackplot(
        pivot.index,
        *values,
        labels=[SERIES_LABELS[series] for series in TECHNOLOGIES],
        colors=[TECHNOLOGY_COLORS[series] for series in TECHNOLOGIES],
        alpha=0.9,
    )
    total = np.sum(values, axis=0)
    axis.plot(pivot.index, total, color="#334155", linewidth=1.8, label="Tổng NLTT")
    axis.set_ylabel("Sản lượng (TWh)")
    axis.set_ylim(0, max(total) * 1.12)
    _set_year_axis(axis)
    _style_axis(axis)


def _render_c02(data: pd.DataFrame, target: Path) -> None:
    fig, axis = _new_figure(1, 1)
    _stacked_area(axis, data)
    axis.legend(loc="upper left", ncol=5)
    endpoint = data[data["year"].eq(2021)].set_index("series_key")["metric_value"]
    annotation = "2021: " + "  ·  ".join(
        f"{SERIES_LABELS[series]} {_vi_number(endpoint[series], 1)}"
        for series in TECHNOLOGIES
    ) + " TWh"
    axis.text(0.99, 0.03, annotation, transform=axis.transAxes, ha="right", fontsize=9, color="#334155")
    fig.subplots_adjust(left=0.09, right=0.96, top=0.84, bottom=0.15)
    _decorate_figure(fig, "C02")
    _save_figure(fig, target)


def _render_c03(data: pd.DataFrame, target: Path) -> None:
    generation = data[data["metric_name"].eq("generation_twh")]
    shares = data[data["metric_name"].eq("technology_share")]
    gen_pivot = generation.pivot(index="subject_key", columns="series_key", values="metric_value").fillna(0)
    share_pivot = shares.pivot(index="subject_key", columns="series_key", values="metric_value")
    order = gen_pivot[TECHNOLOGIES].sum(axis=1).sort_values(ascending=True).index.tolist()
    fig, axes = _new_figure(1, 2, gridspec_kw={"width_ratios": [1.35, 1]})
    left = np.zeros(len(order))
    y = np.arange(len(order))
    for series in TECHNOLOGIES:
        values = gen_pivot.loc[order, series].to_numpy()
        axes[0].barh(y, values, left=left, color=TECHNOLOGY_COLORS[series], label=SERIES_LABELS[series], height=0.68)
        left += values
    axes[0].set_yticks(y, [ENTITY_LABELS[name] for name in order])
    axes[0].set_xlabel("Sản lượng năm 2021 (TWh)")
    axes[0].legend(loc="lower right", fontsize=8)
    _style_axis(axes[0], y_grid=False)
    axes[0].grid(True, axis="x", alpha=0.7)
    for index, total in enumerate(left):
        axes[0].text(total + max(left) * 0.008, index, _vi_number(total, 1), va="center", fontsize=8)

    heat = share_pivot.loc[order, TECHNOLOGIES]
    annotations = heat.map(_format_share_label)
    heatmap_axis = sns.heatmap(
        heat,
        ax=axes[1],
        cmap="YlGnBu",
        vmin=0,
        vmax=1,
        annot=annotations,
        fmt="",
        linewidths=0.6,
        linecolor="white",
        cbar_kws={"label": "Tỷ trọng trong tổng NLTT"},
    )
    colorbar = heatmap_axis.collections[0].colorbar
    if colorbar is None:
        raise Phase4Error("C03 heatmap did not create the required share colorbar")
    colorbar.formatter = FuncFormatter(lambda value, _: _format_share_label(value))
    colorbar.update_ticks()
    axes[1].set_yticklabels([ENTITY_LABELS[name] for name in order], rotation=0)
    # barh places the first (smallest) country at the bottom whereas heatmap
    # places its first row at the top; invert the latter to keep visual order.
    axes[1].invert_yaxis()
    if share_pivot.loc["Brunei", TECHNOLOGIES].isna().all():
        axes[1].text(
            2,
            0.5,
            "Không xác định (tổng = 0)",
            ha="center",
            va="center",
            fontsize=7.5,
            color="#475569",
        )
    axes[1].set_xticklabels([SERIES_LABELS[series] for series in TECHNOLOGIES], rotation=25, ha="right")
    axes[1].set_xlabel("")
    axes[1].set_ylabel("")
    axes[0].set_title("Quy mô tuyệt đối")
    axes[1].set_title("Cơ cấu trong từng nước")
    fig.subplots_adjust(left=0.09, right=0.96, top=0.82, bottom=0.17, wspace=0.34)
    _decorate_figure(fig, "C03")
    _save_figure(fig, target)


def _render_c04(data: pd.DataFrame, target: Path) -> None:
    ordered = data.sort_values(["total_rank", "entity"]).reset_index(drop=True)
    fig, axis = _new_figure(1, 1)
    y = np.arange(len(ordered))
    for index, row in ordered.iterrows():
        color = COUNTRY_COLORS[row["entity"]]
        axis.plot([row["total_rank"], row["modern_rank"]], [index, index], color=color, linewidth=2.2, alpha=0.8)
        axis.scatter(row["total_rank"], index, color=color, s=72, marker="o", zorder=3)
        axis.scatter(row["modern_rank"], index, color=color, s=80, marker="D", edgecolor="white", linewidth=0.8, zorder=4)
        gain = int(row["modern_rank_gain"])
        axis.text(10.28, index, f"{gain:+d}", va="center", ha="left", fontsize=9, color=color, fontweight="bold")
    axis.set_yticks(y, [ENTITY_LABELS[name] for name in ordered["entity"]])
    axis.invert_yaxis()
    axis.set_xlim(0.5, 10.8)
    axis.set_xticks(range(1, 11))
    axis.set_xlabel("Hạng trong ASEAN-10 (1 = cao nhất)")
    axis.set_title("Mỗi đường nối cùng một quốc gia; số bên phải là mức tăng/giảm hạng")
    axis.legend(
        handles=[
            Line2D([], [], marker="o", linestyle="", color="#475569", label="Tổng NLTT", markersize=8),
            Line2D([], [], marker="D", linestyle="", color="#475569", label="Gió + mặt trời", markersize=8),
        ],
        loc="lower left",
    )
    _style_axis(axis, y_grid=False)
    axis.grid(True, axis="x", alpha=0.7)
    fig.subplots_adjust(left=0.16, right=0.93, top=0.82, bottom=0.16)
    _decorate_figure(fig, "C04")
    _save_figure(fig, target)


def _configure_small_multiple_y_axis(
    axis: plt.Axes,
    series: pd.DataFrame,
    *,
    shared: bool,
    global_max: float,
) -> None:
    """Apply the locked y-axis policy, including an explicit all-zero state."""
    if shared:
        axis.set_ylim(0, global_max)
    else:
        local_max = float(series["metric_value"].max())
        if local_max <= NUMERIC_TOLERANCE:
            axis.set_ylim(0, 1)
            axis.set_yticks([0])
            axis.yaxis.set_major_formatter(
                FuncFormatter(lambda value, _: "0" if np.isclose(value, 0) else "")
            )
            axis.text(
                0.5,
                0.5,
                "Toàn kỳ = 0 TWh",
                transform=axis.transAxes,
                ha="center",
                va="center",
                fontsize=9,
                color="#64748B",
                fontweight="bold",
            )
            return
        axis.set_ylim(0, local_max * 1.1)
    axis.yaxis.set_major_formatter(
        FuncFormatter(lambda value, _: _vi_number(value, 0) if value >= 10 else _vi_number(value, 1))
    )


def _render_small_multiples(data: pd.DataFrame, target: Path, *, shared: bool) -> None:
    fig, axes = _new_figure(2, 5, sharex=True, sharey=shared)
    ordered = sorted(ASEAN, key=lambda name: (name != "Vietnam", name))
    global_max = data["metric_value"].max() * 1.08
    for axis, entity in zip(axes.flat, ordered, strict=True):
        series = data[data["subject_key"].eq(entity)].sort_values("year")
        axis.plot(series["year"], series["metric_value"], color=COUNTRY_COLORS[entity], linewidth=2.2)
        axis.fill_between(series["year"], series["metric_value"], color=COUNTRY_COLORS[entity], alpha=0.08)
        axis.set_title(ENTITY_LABELS[entity], fontsize=10.5, color=COUNTRY_COLORS[entity])
        axis.set_xlim(YEAR_START, YEAR_END)
        axis.set_xticks([2000, 2010, 2021])
        axis.tick_params(axis="x", rotation=0, labelsize=8)
        _configure_small_multiple_y_axis(
            axis,
            series,
            shared=shared,
            global_max=global_max,
        )
        _style_axis(axis)
    fig.text(0.018, 0.5, "Sản lượng tổng NLTT (TWh)", rotation=90, va="center", fontsize=10)
    chart_id = "C05" if shared else "C06"
    fig.subplots_adjust(left=0.065, right=0.97, top=0.82, bottom=0.16, hspace=0.42, wspace=0.22)
    _decorate_figure(fig, chart_id)
    _save_figure(fig, target)


def _render_c07(data: pd.DataFrame, target: Path) -> None:
    fig, axis = _new_figure(1, 1)
    _stacked_area(axis, data)
    axis.legend(loc="upper left", ncol=5)
    axis.axvspan(2018, 2021, color="#D55E00", alpha=0.045)
    axis.text(2018.15, axis.get_ylim()[1] * 0.86, "Gió và mặt trời tăng rõ\nở cuối cửa sổ", color="#D55E00", fontsize=10, fontweight="bold")
    fig.subplots_adjust(left=0.09, right=0.96, top=0.84, bottom=0.15)
    _decorate_figure(fig, "C07")
    _save_figure(fig, target)


def _render_c08(data: pd.DataFrame, target: Path) -> None:
    fig, axes = _new_figure(2, 2, sharex=True)
    for column, series_key in enumerate(POSITION_SERIES):
        subset = data[data["series_key"].eq(series_key)].sort_values("year")
        color = TECHNOLOGY_COLORS[series_key]
        axes[0, column].plot(subset["year"], subset["generation_twh"], color=color, linewidth=2.5)
        axes[1, column].plot(subset["year"], subset["share_electricity_pct"], color=color, linewidth=2.5)
        axes[0, column].set_title(SERIES_LABELS[series_key])
        for axis in axes[:, column]:
            _set_year_axis(axis)
            _style_axis(axis)
        opposite = subset[subset["direction_relation_text"].eq("OPPOSITE_DIRECTION")]
        for _, row in opposite.iterrows():
            for axis, value in [(axes[0, column], row["generation_twh"]), (axes[1, column], row["share_electricity_pct"])]:
                axis.scatter(row["year"], value, s=70, facecolor="white", edgecolor="#D55E00", linewidth=2, zorder=4)
            axes[0, column].annotate(
                f"{int(row['year'])}: ngược hướng",
                (row["year"], row["generation_twh"]),
                xytext=(0, 10),
                textcoords="offset points",
                ha="center",
                fontsize=8,
                color="#D55E00",
            )
    axes[0, 0].set_ylabel("Sản lượng (TWh)")
    axes[1, 0].set_ylabel("Tỷ trọng trong tổng điện (%)")
    axes[1, 0].set_xlabel("Năm")
    axes[1, 1].set_xlabel("Năm")
    for axis in axes[1, :]:
        axis.yaxis.set_major_formatter(FuncFormatter(lambda value, _: f"{value:.0f}%"))
    fig.subplots_adjust(left=0.085, right=0.96, top=0.82, bottom=0.16, hspace=0.25, wspace=0.18)
    _decorate_figure(fig, "C08")
    _save_figure(fig, target)


def _render_c09(data: pd.DataFrame, target: Path) -> None:
    fig, axes = _new_figure(1, 2, sharex=True)
    time_rows = data[data["record_type"].eq("TIME_SERIES")]
    rank_rows = data[data["record_type"].eq("GLOBAL_RANK")]
    colors = {"Vietnam": "#D55E00", "ASEAN_9_TOTAL": "#64748B"}
    labels = {"Vietnam": "Việt Nam", "ASEAN_9_TOTAL": "Tổng ASEAN-9"}
    for axis, series_key in zip(axes, POSITION_SERIES, strict=True):
        subset = time_rows[time_rows["series_key"].eq(series_key)]
        for subject in ["ASEAN_9_TOTAL", "Vietnam"]:
            line = subset[subset["display_subject"].eq(subject)].sort_values("year")
            axis.plot(line["year"], line["metric_value"], color=colors[subject], linewidth=2.7, label=labels[subject])
            endpoint = line.iloc[-1]
            axis.annotate(
                f"{labels[subject]}: {_vi_number(endpoint['metric_value'], 1)} TWh",
                (endpoint["year"], endpoint["metric_value"]),
                xytext=(-4, 8 if subject == "Vietnam" else -15),
                textcoords="offset points",
                ha="right",
                fontsize=8.5,
                color=colors[subject],
                fontweight="bold" if subject == "Vietnam" else "normal",
                # The ASEAN-9 series rises steeply into its endpoint, so both lines can
                # cross this label. A flat white halo keeps it readable without moving it.
                bbox={"boxstyle": "round,pad=0.18", "facecolor": "white", "edgecolor": "none", "alpha": 0.85},
            )
        rank = rank_rows[rank_rows["series_key"].eq(series_key)].iloc[0]
        axis.text(
            0.03,
            0.84,
            f"Việt Nam năm 2021: hạng {int(rank['metric_value'])}/{int(rank['ranked_n'])} toàn cầu",
            transform=axis.transAxes,
            va="top",
            fontsize=9,
            color="#D55E00",
            fontweight="bold",
            bbox={"boxstyle": "round,pad=0.35", "facecolor": "#FFF7ED", "edgecolor": "#FED7AA"},
        )
        axis.set_title(SERIES_LABELS[series_key])
        axis.set_ylabel("Sản lượng (TWh)")
        axis.set_ylim(bottom=0)
        _set_year_axis(axis)
        _style_axis(axis)
    axes[0].legend(loc="upper left", bbox_to_anchor=(0, 0.84))
    fig.subplots_adjust(left=0.08, right=0.96, top=0.82, bottom=0.16, wspace=0.18)
    _decorate_figure(fig, "C09")
    _save_figure(fig, target)


def _render_c10(data: pd.DataFrame, target: Path) -> None:
    values = data[data["record_type"].eq("COUNTRY_VALUE")].copy()
    fig, axes = _new_figure(1, 2, gridspec_kw={"width_ratios": [1, 1.35]})
    y = np.arange(len(TECHNOLOGIES))
    classes = ["Zero", "Trace", "Material", "Missing"]
    class_labels = {
        "Zero": "Bằng 0",
        "Trace": "Dương dưới ngưỡng (0–0,05)",
        "Material": "Đạt ngưỡng (≥0,05)",
        "Missing": "Thiếu",
    }
    class_colors = {"Zero": "#CBD5E1", "Trace": "#93C5FD", "Material": "#2563EB", "Missing": "#F8FAFC"}
    left = np.zeros(len(TECHNOLOGIES))
    for value_class in classes:
        counts = np.array([
            int(((values["series_key"].eq(series)) & values["display_class"].eq(value_class)).sum())
            for series in TECHNOLOGIES
        ])
        axes[0].barh(y, counts, left=left, color=class_colors[value_class], edgecolor="#94A3B8", linewidth=0.4, label=class_labels[value_class])
        for index, (start, count) in enumerate(zip(left, counts, strict=True)):
            if count >= 15:
                axes[0].text(start + count / 2, index, str(count), ha="center", va="center", fontsize=8, color="#0F172A")
        left += counts
    axes[0].set_yticks(y, [SERIES_LABELS[series] for series in TECHNOLOGIES])
    axes[0].invert_yaxis()
    axes[0].set_xlim(0, 215)
    axes[0].set_xlabel("Số quốc gia/vùng lãnh thổ (mẫu cố định 215)")
    axes[0].set_title("Khối lượng tại 0 và ngưỡng quan sát")
    axes[0].legend(loc="lower center", bbox_to_anchor=(0.5, -0.27), ncol=2, fontsize=8)
    _style_axis(axes[0], y_grid=False)
    axes[0].grid(True, axis="x", alpha=0.7)

    positive = values[pd.to_numeric(values["metric_value"], errors="coerce").gt(0)].copy()
    positive["technology_label"] = positive["series_key"].map(SERIES_LABELS)
    order = [SERIES_LABELS[series] for series in TECHNOLOGIES]
    palette = {SERIES_LABELS[series]: TECHNOLOGY_COLORS[series] for series in TECHNOLOGIES}
    sns.boxplot(data=positive, x="metric_value", y="technology_label", order=order, palette=palette, hue="technology_label", legend=False, showfliers=False, whis=(10, 90), ax=axes[1])
    sns.stripplot(data=positive, x="metric_value", y="technology_label", order=order, palette=palette, hue="technology_label", legend=False, size=2.3, jitter=0.22, alpha=0.35, ax=axes[1])
    axes[1].set_xscale("log")
    axes[1].set_xlabel("Sản lượng dương (TWh, thang logarit)")
    axes[1].set_ylabel("")
    axes[1].set_title("Phân phối dương vẫn lệch phải mạnh")
    _style_axis(axes[1], y_grid=False)
    axes[1].grid(True, axis="x", which="both", alpha=0.5)
    fig.subplots_adjust(left=0.09, right=0.96, top=0.82, bottom=0.25, wspace=0.33)
    _decorate_figure(fig, "C10")
    _save_figure(fig, target)


def _render_c11(data: pd.DataFrame, target: Path) -> None:
    fig, axes = _new_figure(2, 2, sharex=True, sharey=True)
    statuses = {
        "MATERIAL_AT_WINDOW_START": (2000, "<", "Đã đạt trước/đúng đầu cửa sổ"),
        "OBSERVED_MATERIAL_TRANSITION": (None, "o", "Quan sát được năm chuyển"),
        "NOT_MATERIAL_BY_2021": (2021, ">", "Chưa đạt đến cuối cửa sổ"),
    }
    ordered_entities = sorted(ASEAN, key=lambda name: (name != "Vietnam", name))
    for axis, series_key in zip(axes.flat, TECHNOLOGIES, strict=True):
        subset = data[data["series_key"].eq(series_key)].set_index("entity")
        color = TECHNOLOGY_COLORS[series_key]
        for index, entity in enumerate(ordered_entities):
            row = subset.loc[entity]
            status = row["material_entry_status_text"]
            fixed_year, marker, _ = statuses[status]
            year = float(row["observed_material_entry_year"]) if status == "OBSERVED_MATERIAL_TRANSITION" else fixed_year
            face = color if status == "OBSERVED_MATERIAL_TRANSITION" else "white"
            axis.scatter(year, index, marker=marker, s=65, facecolor=face, edgecolor=color, linewidth=1.6, zorder=3)
        axis.set_title(SERIES_LABELS[series_key], color=color)
        axis.set_xlim(1999.2, 2021.8)
        axis.set_xticks([2000, 2005, 2010, 2015, 2021])
        axis.set_yticks(range(len(ordered_entities)), [ENTITY_LABELS[name] for name in ordered_entities])
        axis.invert_yaxis()
        axis.grid(True, axis="x", alpha=0.65)
        axis.grid(False, axis="y")
        axis.spines[["top", "right"]].set_visible(False)
    handles = [
        Line2D([], [], marker=marker, linestyle="", markerfacecolor="white" if status != "OBSERVED_MATERIAL_TRANSITION" else "#64748B", markeredgecolor="#64748B", label=label, markersize=8)
        for status, (_, marker, label) in statuses.items()
    ]
    fig.legend(handles=handles, loc="lower center", bbox_to_anchor=(0.5, 0.105), ncol=3, fontsize=8.5)
    fig.subplots_adjust(left=0.12, right=0.96, top=0.82, bottom=0.21, hspace=0.32, wspace=0.16)
    _decorate_figure(fig, "C11")
    _save_figure(fig, target)


def _render_c12(concentration: pd.DataFrame, growth: pd.DataFrame, target: Path) -> None:
    fig, axes = _new_figure(1, 2, gridspec_kw={"width_ratios": [1.1, 1]})
    top3 = concentration[concentration["metric_name"].eq("top3_share")]
    colors = {"RENEWABLE_TOTAL": "#374151", "MODERN_WIND_SOLAR": "#D55E00"}
    for series_key in POSITION_SERIES:
        series = top3[top3["series_key"].eq(series_key)].sort_values("year")
        axes[0].plot(series["year"], series["metric_value"] * 100, color=colors[series_key], linewidth=2.7, label=SERIES_LABELS[series_key])
        end = series.iloc[-1]
        axes[0].annotate(_format_share_label(end["metric_value"]), (end["year"], end["metric_value"] * 100), xytext=(-4, 7), textcoords="offset points", ha="right", color=colors[series_key], fontsize=9, fontweight="bold")
    axes[0].set_ylabel("Tỷ trọng của ba nước dẫn đầu")
    axes[0].yaxis.set_major_formatter(PercentFormatter(100, decimals=0))
    # Leave headroom so early years at exactly 100% are not clipped by the top spine.
    axes[0].set_ylim(0, 104)
    axes[0].set_title("Mức sản lượng do ba nước dẫn đầu nắm giữ")
    axes[0].legend(loc="upper left")
    _set_year_axis(axes[0])
    _style_axis(axes[0])

    bars = growth[(growth["subject_type"].eq("COUNTRY")) & growth["metric_name"].eq("positive_growth_share")].sort_values("metric_value", ascending=True)
    labels = [ENTITY_LABELS[name] for name in bars["subject_key"]]
    bar_colors = [COUNTRY_COLORS[name] for name in bars["subject_key"]]
    axes[1].barh(labels, bars["metric_value"] * 100, color=bar_colors, height=0.68)
    for index, value in enumerate(bars["metric_value"] * 100):
        if value >= 1:
            axes[1].text(value + 0.7, index, f"{_vi_number(value, 1)}%", va="center", fontsize=8.5)
    aggregate = growth[growth["subject_type"].eq("REGION")].set_index("metric_name")["metric_value"]
    axes[1].text(
        0.98,
        0.04,
        f"Tổng phần tăng dương: {_vi_number(aggregate['gross_positive_addition_twh'], 1)} TWh\nBa nước đóng góp nhiều nhất: {_format_share_label(aggregate['top3_positive_growth_share'])}",
        transform=axes[1].transAxes,
        ha="right",
        fontsize=8.5,
        color="#475569",
    )
    axes[1].set_xlabel("Tỷ trọng trong tổng phần tăng dương gió + mặt trời")
    axes[1].xaxis.set_major_formatter(PercentFormatter(100, decimals=0))
    axes[1].set_title("Đóng góp tăng dương, 2010–2021")
    _style_axis(axes[1], y_grid=False)
    axes[1].grid(True, axis="x", alpha=0.7)
    fig.subplots_adjust(left=0.09, right=0.96, top=0.82, bottom=0.17, wspace=0.32)
    _decorate_figure(fig, "C12")
    _save_figure(fig, target)


def _render_c13(data: pd.DataFrame, target: Path) -> None:
    fig, axes = _new_figure(1, 2, sharey=True)
    present = {technology: sorted(data[data["technology"].eq(technology)]["entity"].unique()) for technology in ["WIND", "SOLAR"]}
    for axis, technology in zip(axes, ["WIND", "SOLAR"], strict=True):
        subset = data[data["technology"].eq(technology)]
        for entity in present[technology]:
            series = subset[subset["entity"].eq(entity)].sort_values("year")
            axis.plot(series["year"], series["capacity_factor_ratio"] * 100, color=COUNTRY_COLORS[entity], linewidth=2.2, label=ENTITY_LABELS[entity])
            end = series.iloc[-1]
            axis.scatter(end["year"], end["capacity_factor_ratio"] * 100, color=COUNTRY_COLORS[entity], s=30)
        axis.set_title(f"{SERIES_LABELS[technology]} — {len(present[technology])} nước có dữ liệu")
        axis.set_ylabel("Hệ số sử dụng công suất xấp xỉ")
        axis.yaxis.set_major_formatter(PercentFormatter(100, decimals=0))
        axis.set_ylim(0, max(40, float(data["capacity_factor_ratio"].max() * 100 * 1.18)))
        _set_year_axis(axis)
        _style_axis(axis)
        axis.legend(loc="upper left", fontsize=8.5)
    missing_wind = [ENTITY_LABELS[name] for name in ASEAN if name not in present["WIND"]]
    missing_solar = [ENTITY_LABELS[name] for name in ASEAN if name not in present["SOLAR"]]
    footer = (
        "Nguồn: output/clean/fact_energy.csv. Phạm vi: ASEAN-10, 2000–2021; chỉ giữ cặp năm–nước có sản lượng quan sát được và công suất > 0. "
        "Công thức: TWh / (GW × 8,76). Công suất là số cuối năm, vì vậy đây là chỉ báo xấp xỉ, không phải hiệu quả kỹ thuật. "
        f"Thiếu công suất gió: {', '.join(missing_wind)}. Thiếu công suất mặt trời: {', '.join(missing_solar)}."
    )
    fig.subplots_adjust(left=0.08, right=0.96, top=0.82, bottom=0.22, wspace=0.18)
    _decorate_figure(fig, "C13", footer)
    _save_figure(fig, target)


def render_charts(data: dict[str, pd.DataFrame], target_dir: Path) -> None:
    _configure_plotting()
    renderers = {
        "C01": lambda path: _render_c01(data["global_context"], path),
        "C02": lambda path: _render_c02(data["asean_technology"], path),
        "C03": lambda path: _render_c03(data["country_mix_2021"], path),
        "C04": lambda path: _render_c04(data["rank_shift_2021"], path),
        "C05": lambda path: _render_small_multiples(data["small_multiples"], path, shared=True),
        "C06": lambda path: _render_small_multiples(data["small_multiples"], path, shared=False),
        "C07": lambda path: _render_c07(data["vietnam_technology"], path),
        "C08": lambda path: _render_c08(data["vietnam_q5"], path),
        "C09": lambda path: _render_c09(data["vietnam_position"], path),
        "C10": lambda path: _render_c10(data["global_distribution"], path),
        "C11": lambda path: _render_c11(data["entry_lag"], path),
        "C12": lambda path: _render_c12(data["concentration"], data["growth_contribution"], path),
        "C13": lambda path: _render_c13(data["capacity_factor"], path),
    }
    target_dir.mkdir(parents=True, exist_ok=True)
    for spec in CHART_SPECS:
        renderers[spec["chart_id"]](target_dir / spec["file_name"])


def _close(actual: float, expected: float, tolerance: float = NUMERIC_TOLERANCE) -> bool:
    return bool(np.isclose(float(actual), float(expected), rtol=0.0, atol=tolerance))


def _source_has_forbidden_plot_calls() -> list[str]:
    """Return prohibited plotting calls found in this source via the Python AST."""
    tree = ast.parse(Path(__file__).read_text(encoding="utf-8"))
    prohibited = {"pie", "twinx", "twiny"}
    found: list[str] = []
    for node in ast.walk(tree):
        if not isinstance(node, ast.Call):
            continue
        name = ""
        if isinstance(node.func, ast.Attribute):
            name = node.func.attr
        elif isinstance(node.func, ast.Name):
            name = node.func.id
        if name.lower() in prohibited:
            found.append(name.lower())
    return sorted(set(found))


def _data_fact_keys(data: dict[str, pd.DataFrame]) -> set[str]:
    keys: set[str] = set()
    for frame in data.values():
        for column in frame.columns:
            if column == "fact_key" or column.endswith("_fact_key"):
                keys.update(
                    value
                    for value in frame[column].dropna().astype(str)
                    if value.startswith("[")
                )
    return keys


def _same_source_value(actual: Any, expected: Any) -> bool:
    actual_missing = bool(pd.isna(actual)) if not isinstance(actual, (list, dict)) else False
    expected_missing = bool(pd.isna(expected)) if not isinstance(expected, (list, dict)) else False
    if actual_missing or expected_missing:
        return actual_missing and expected_missing
    if isinstance(actual, (int, float, np.number)) and isinstance(expected, (int, float, np.number)):
        return _close(float(actual), float(expected))
    return str(actual) == str(expected)


def _traceability_result(inputs: Inputs, data: dict[str, pd.DataFrame]) -> tuple[bool, dict[str, Any]]:
    facts = inputs.facts
    duplicate_source_keys = int(facts.duplicated("fact_key", keep=False).sum())
    lookup = facts.drop_duplicates("fact_key").set_index("fact_key")
    mismatches: list[str] = []
    checked = 0
    for sheet_name, frame in data.items():
        for key_column in [column for column in frame.columns if column == "fact_key" or column.endswith("_fact_key")]:
            base = "" if key_column == "fact_key" else key_column[: -len("_fact_key")]
            for row_index, row in frame.iterrows():
                key = row.get(key_column)
                if pd.isna(key) or not str(key).startswith("["):
                    continue
                checked += 1
                key = str(key)
                if key not in lookup.index:
                    mismatches.append(f"{sheet_name}:{row_index}:{key_column}:missing_key")
                    continue
                source = lookup.loc[key]
                comparisons: list[tuple[str, str]] = []
                if key_column == "fact_key":
                    comparisons = [
                        ("metric_value", "metric_value"),
                        ("metric_text_value", "metric_text_value"),
                        ("unit", "unit"),
                        ("value_status", "value_status"),
                        ("exclusion_reason", "exclusion_reason"),
                    ]
                else:
                    comparisons = [
                        (base, "metric_value"),
                        (f"{base}_text", "metric_text_value"),
                        (f"{base}_unit", "unit"),
                        (f"{base}_status", "value_status"),
                        (f"{base}_value_status", "value_status"),
                        (f"{base}_exclusion_reason", "exclusion_reason"),
                    ]
                for output_column, source_column in comparisons:
                    if output_column in frame.columns and source_column in source.index:
                        if not _same_source_value(row[output_column], source[source_column]):
                            mismatches.append(
                                f"{sheet_name}:{row_index}:{output_column}!={source_column}"
                            )
                if "source_method_version" in frame.columns and not _same_source_value(
                    row["source_method_version"], source["method_version"]
                ):
                    mismatches.append(f"{sheet_name}:{row_index}:source_method_version")
    evidence = {
        "checked_key_occurrences": checked,
        "unique_used_keys": len(_data_fact_keys(data)),
        "duplicate_authoritative_keys": duplicate_source_keys,
        "mismatch_count": len(mismatches),
        "first_mismatches": mismatches[:10],
    }
    return duplicate_source_keys == 0 and not mismatches, evidence


def _all_years_in_domain(data: dict[str, pd.DataFrame]) -> tuple[bool, list[int]]:
    years: set[int] = set()
    valid = True
    for frame in data.values():
        if "year" not in frame.columns:
            continue
        numeric = pd.to_numeric(frame["year"], errors="coerce").dropna()
        years.update(int(value) for value in numeric)
        valid = valid and bool(numeric.between(YEAR_START, YEAR_END).all())
    return valid, sorted(years)


def _no_unexplained_nonfinite(data: dict[str, pd.DataFrame]) -> tuple[bool, int]:
    """Reject infinities; structural/typed blanks remain represented by status fields."""
    infinite_count = 0
    for name, frame in data.items():
        if name in {"manifest", "metadata", "method_notes"}:
            continue
        numeric = frame.select_dtypes(include=[np.number])
        if not numeric.empty:
            infinite_count += int(np.isinf(numeric.to_numpy(dtype=float, na_value=np.nan)).sum())
    unexplained_missing = 0
    literal_nan_strings = 0
    for frame in data.values():
        object_frame = frame.select_dtypes(include=["object", "string"])
        if not object_frame.empty:
            literal_nan_strings += int(
                object_frame.apply(
                    lambda column: column.notna()
                    & column.astype(str).str.strip().str.lower().eq("nan")
                ).to_numpy().sum()
            )
        if "metric_value" in frame.columns:
            missing = frame["metric_value"].isna()
            if missing.any():
                text_explains = (
                    frame["metric_text_value"].notna()
                    if "metric_text_value" in frame.columns
                    else pd.Series(False, index=frame.index)
                )
                status_explains = (
                    frame["value_status"].fillna("").astype(str).ne("")
                    & ~frame["value_status"].fillna("").astype(str).isin(["OBSERVED", "CALCULATED"])
                    if "value_status" in frame.columns
                    else pd.Series(False, index=frame.index)
                )
                unexplained_missing += int((missing & ~text_explains & ~status_explains).sum())
        for base in ["generation_twh", "share_electricity_pct", "direction_relation", "observed_material_entry_year"]:
            if base not in frame.columns:
                continue
            missing = frame[base].isna()
            text_column = f"{base}_text"
            status_column = f"{base}_status"
            value_status_column = f"{base}_value_status"
            text_explains = frame[text_column].notna() if text_column in frame else pd.Series(False, index=frame.index)
            status_explains = pd.Series(False, index=frame.index)
            if status_column in frame:
                status_explains |= frame[status_column].fillna("").astype(str).ne("")
            if value_status_column in frame:
                status_explains |= frame[value_status_column].fillna("").astype(str).ne("")
            # Long-form fact sheets already handled metric_value above.
            if base != "generation_twh" or "metric_value" not in frame.columns:
                unexplained_missing += int((missing & ~text_explains & ~status_explains).sum())
    invalid_count = infinite_count + unexplained_missing + literal_nan_strings
    return invalid_count == 0, invalid_count


def _sheet_contract_result(data: dict[str, pd.DataFrame]) -> tuple[bool, dict[str, Any]]:
    details: dict[str, Any] = {}
    valid = list(data) == list(CHART_DATA_SHEETS)
    for name in CHART_DATA_SHEETS:
        frame = data.get(name)
        if frame is None:
            details[name] = "MISSING"
            valid = False
            continue
        missing = sorted(set(SHEET_REQUIRED_COLUMNS[name]) - set(frame.columns))
        method_values = sorted(frame["method_version"].dropna().astype(str).unique()) if "method_version" in frame else []
        fingerprint_values = sorted(frame["phase2_input_fingerprint"].dropna().astype(str).unique()) if "phase2_input_fingerprint" in frame else []
        provenance_ok = method_values == [METHOD_VERSION] and fingerprint_values == [EXPECTED_PHASE2_FINGERPRINT]
        details[name] = {
            "rows": int(len(frame)),
            "missing_columns": missing,
            "method_versions": method_values,
            "fingerprints": fingerprint_values,
        }
        valid = valid and not missing and provenance_ok
    return valid, details


def _unique_key_result(data: dict[str, pd.DataFrame]) -> tuple[bool, dict[str, int]]:
    duplicates: dict[str, int] = {}
    valid = True
    for name, keys in SHEET_KEY_COLUMNS.items():
        frame = data[name]
        missing = set(keys) - set(frame.columns)
        count = len(frame) if missing else int(frame.duplicated(list(keys), keep=False).sum())
        duplicates[name] = count
        valid = valid and count == 0
    return valid, duplicates


def _phase3_total_lookup(inputs: Inputs, subject_key: Any) -> pd.DataFrame:
    return _period_to_year(
        select_facts(
            inputs.facts,
            subject_type="COUNTRY" if subject_key != "ASEAN_10" else "REGION",
            subject_key=subject_key,
            period_type="YEAR",
            series_key="RENEWABLE_TOTAL",
            metric_name="generation_twh",
            comparison_scope="ASEAN_10",
            filter_variant="NOT_APPLICABLE",
        )
    )[["subject_key", "year", "metric_value"]]


def _validate_chart_contracts(
    collector: ValidationCollector,
    inputs: Inputs,
    data: dict[str, pd.DataFrame],
) -> None:
    # C01: 22 annual values for each of two levels and one ratio, plus locked endpoints.
    c01 = data["global_context"]
    endpoint_values: dict[str, dict[int, float]] = {}
    for label, expected_years in EXPECTED_C01_ENDPOINTS.items():
        if label == "ASEAN_SHARE_WORLD":
            selected = c01[c01["metric_name"].eq("asean_share_world")]
        else:
            selected = c01[
                c01["subject_key"].eq(label) & c01["metric_name"].eq("generation_twh")
            ]
        endpoint_values[label] = {
            year: float(selected.loc[selected["year"].eq(year), "metric_value"].iloc[0])
            for year in expected_years
        }
    c01_endpoints_ok = all(
        (
            _close(endpoint_values[label][year], expected)
            if label == "ASEAN_SHARE_WORLD"
            else round(endpoint_values[label][year], 6) == round(expected, 6)
        )
        for label, expected_years in EXPECTED_C01_ENDPOINTS.items()
        for year, expected in expected_years.items()
    )
    c01_positive = c01.loc[c01["metric_name"].eq("generation_twh"), "metric_value"].gt(0).all()
    collector.add(
        "P4_C01_CONTRACT", "CHART_CONTRACT", "C01",
        len(c01) == 66 and c01["year"].nunique() == 22 and c01_endpoints_ok and c01_positive,
        "66 values; 22 years; locked endpoints; log inputs > 0",
        {"rows": len(c01), "years": c01["year"].nunique(), "endpoints": endpoint_values},
        "Separate panels prevent a dual axis; only positive TWh values enter the logarithmic panel.",
    )

    # C02: components are complete and reconcile to the Phase 3 regional total.
    c02 = data["asean_technology"]
    c02_sum = c02.groupby("year", as_index=False)["metric_value"].sum()
    c02_total = _phase3_total_lookup(inputs, "ASEAN_10")[["year", "metric_value"]].rename(columns={"metric_value": "total"})
    c02_check = c02_sum.merge(c02_total, on="year", validate="one_to_one")
    c02_diff = float((c02_check["metric_value"] - c02_check["total"]).abs().max())
    c02_2021 = c02[c02["year"].eq(2021)].set_index("series_key")["metric_value"].to_dict()
    c02_locked = all(
        round(float(c02_2021[key]), 6) == round(float(value), 6)
        for key, value in EXPECTED_C02_2021_TWH.items()
    )
    collector.add(
        "P4_C02_CONTRACT", "CHART_CONTRACT", "C02",
        len(c02) == 88 and c02["year"].nunique() == 22 and c02["series_key"].nunique() == 4 and c02_diff <= NUMERIC_TOLERANCE and c02_locked,
        "88 values; four complete components; annual sum difference <= 1e-9; locked 2021 values",
        {"rows": len(c02), "max_sum_difference": c02_diff, "values_2021": c02_2021},
        "Component reconciliation prevents a total from being reconstructed from incomplete data.",
    )

    # C03: two representations of ten-country, four-technology 2021 data.
    c03 = data["country_mix_2021"]
    c03_generation = c03[c03["metric_name"].eq("generation_twh")]
    c03_sums = c03_generation.groupby("entity", as_index=False)["metric_value"].sum()
    c03_totals = _phase3_total_lookup(inputs, ASEAN)
    c03_totals = c03_totals[c03_totals["year"].eq(2021)][["subject_key", "metric_value"]].rename(columns={"subject_key": "entity", "metric_value": "total"})
    c03_recon = c03_sums.merge(c03_totals, on="entity", validate="one_to_one")
    c03_diff = float((c03_recon["metric_value"] - c03_recon["total"]).abs().max())
    c03_undefined = c03[c03["metric_value"].isna()]
    c03_positive_labels = c03.loc[
        c03["metric_name"].eq("technology_share") & c03["metric_value"].gt(0),
        "metric_value",
    ].map(_format_share_label)
    c03_label_ok = bool(
        c03_positive_labels.ne("").all()
        and ~c03_positive_labels.isin({"0%", "0,0%"}).any()
    )
    c03_status_ok = (
        len(c03_undefined) == 4
        and set(c03_undefined["entity"]) == {"Brunei"}
        and c03_undefined["value_status"].eq("NOT_DEFINED").all()
        and c03_undefined["exclusion_reason"].eq("ZERO_DENOMINATOR").all()
    )
    collector.add(
        "P4_C03_CONTRACT", "CHART_CONTRACT", "C03",
        len(c03) == 80 and set(c03["entity"]) == set(ASEAN) and c03["series_key"].nunique() == 4 and c03["metric_name"].nunique() == 2 and c03_diff <= NUMERIC_TOLERANCE and c03_status_ok and c03_label_ok,
        "80 records: 76 numeric and four Brunei shares explicitly undefined; stacked total reconciles; no positive share is labelled zero",
        {"rows": len(c03), "numeric_values": int(c03["metric_value"].notna().sum()), "undefined_values": len(c03_undefined), "countries": c03["entity"].nunique(), "max_sum_difference": c03_diff, "positive_share_labels_nonzero": c03_label_ok},
        "Country rows exclude aggregates and use identical ordering in both panels.",
    )

    c04 = data["rank_shift_2021"]
    c04_actual = c04.set_index("entity")["modern_rank_gain"].to_dict()
    c04_shifts_ok = all(_close(c04_actual[key], value) for key, value in EXPECTED_C04_RANK_SHIFTS.items())
    collector.add(
        "P4_C04_CONTRACT", "CHART_CONTRACT", "C04",
        len(c04) == 10 and set(c04["entity"]) == set(ASEAN) and c04_shifts_ok,
        EXPECTED_C04_RANK_SHIFTS,
        {key: c04_actual[key] for key in EXPECTED_C04_RANK_SHIFTS},
        "Rank values come directly from Phase 3, preserving ties; entity colors do not encode rank.",
    )

    small = data["small_multiples"]
    small_ok = len(small) == 220 and set(small["entity"]) == set(ASEAN) and set(small["year"].astype(int)) == set(YEARS)
    collector.add(
        "P4_C05_CONTRACT", "CHART_CONTRACT", "C05", small_ok,
        "220 identical observations; ASEAN-10; 2000-2021; shared y limit",
        {"rows": len(small), "countries": small["entity"].nunique(), "years": small["year"].nunique()},
        "All ten panels use the same explicit y limit derived from this single data sheet.",
    )
    collector.add(
        "P4_C06_CONTRACT", "CHART_CONTRACT", "C06", small_ok,
        "same 220 observations as C05; independent y limits",
        {"rows": len(small), "data_sheet": "small_multiples"},
        "Only the y-scale policy changes from C05; data and ordering remain identical.",
    )

    c07 = data["vietnam_technology"]
    c07_sum = c07.groupby("year", as_index=False)["metric_value"].sum()
    c07_total = _phase3_total_lookup(inputs, "Vietnam")[["year", "metric_value"]].rename(columns={"metric_value": "total"})
    c07_recon = c07_sum.merge(c07_total, on="year", validate="one_to_one")
    c07_diff = float((c07_recon["metric_value"] - c07_recon["total"]).abs().max())
    collector.add(
        "P4_C07_CONTRACT", "CHART_CONTRACT", "C07",
        len(c07) == 88 and c07["year"].nunique() == 22 and c07["series_key"].nunique() == 4 and c07_diff <= NUMERIC_TOLERANCE,
        "88 values; four technologies x 22 years; component sum difference <= 1e-9",
        {"rows": len(c07), "max_sum_difference": c07_diff},
        "The chart describes composition and trend only; it does not infer causes.",
    )

    c08 = data["vietnam_q5"]
    opposite_actual = {
        series: tuple(sorted(c08.loc[c08["series_key"].eq(series) & c08["direction_relation_text"].eq("OPPOSITE_DIRECTION"), "year"].astype(int)))
        for series in POSITION_SERIES
    }
    collector.add(
        "P4_C08_CONTRACT", "CHART_CONTRACT", "C08",
        len(c08) == 44 and c08["generation_twh"].notna().sum() + c08["share_electricity_pct"].notna().sum() == 88 and opposite_actual == EXPECTED_C08_OPPOSITE_DIRECTION_YEARS,
        EXPECTED_C08_OPPOSITE_DIRECTION_YEARS,
        {"paired_rows": len(c08), "numeric_values": int(c08["generation_twh"].notna().sum() + c08["share_electricity_pct"].notna().sum()), "opposite_years": opposite_actual},
        "Separate panels compare TWh and percentage without a dual axis.",
    )

    c09 = data["vietnam_position"]
    c09_time = c09[c09["record_type"].eq("TIME_SERIES")]
    c09_rank = c09[c09["record_type"].eq("GLOBAL_RANK")]
    values_2021: dict[str, dict[str, float]] = {}
    ranks_2021: dict[str, dict[str, int]] = {}
    for series in POSITION_SERIES:
        values_2021[series] = {}
        for subject in ["Vietnam", "ASEAN_9"]:
            display = "Vietnam" if subject == "Vietnam" else "ASEAN_9_TOTAL"
            values_2021[series][subject] = float(c09_time.loc[c09_time["series_key"].eq(series) & c09_time["display_subject"].eq(display) & c09_time["year"].eq(2021), "metric_value"].iloc[0])
        rank_row = c09_rank[c09_rank["series_key"].eq(series)].iloc[0]
        ranks_2021[series] = {"rank": int(rank_row["metric_value"]), "ranked_n": int(rank_row["ranked_n"])}
    c09_values_ok = all(
        round(float(values_2021[s][k]), 6) == round(float(v), 6)
        for s, groups in EXPECTED_C09_2021_TWH.items()
        for k, v in groups.items()
    )
    collector.add(
        "P4_C09_CONTRACT", "CHART_CONTRACT", "C09",
        len(c09_time) == 88 and len(c09_rank) == 2 and c09_values_ok and ranks_2021 == EXPECTED_C09_2021_GLOBAL_RANKS,
        {"values_2021": EXPECTED_C09_2021_TWH, "ranks_2021": EXPECTED_C09_2021_GLOBAL_RANKS},
        {"values_2021": values_2021, "ranks_2021": ranks_2021},
        "ASEAN-9 is a sum, not a substituted average; global rank is annotated only at 2021.",
    )

    c10 = data["global_distribution"]
    c10_values = c10[c10["record_type"].eq("COUNTRY_VALUE")]
    c10_summary = c10[c10["record_type"].eq("SUMMARY")].set_index("series_key")["metric_value"].to_dict()
    c10_zeros = c10_values[c10_values["display_class"].eq("Zero")].groupby("series_key").size().to_dict()
    observed_actual = {key: int(c10_summary[key]) for key in EXPECTED_C10_OBSERVED_N}
    zero_actual = {key: int(c10_zeros.get(key, 0)) for key in EXPECTED_C10_ZERO_N}
    c10_positive = pd.to_numeric(c10_values.loc[c10_values["metric_value"].gt(0), "metric_value"], errors="coerce").gt(0).all()
    collector.add(
        "P4_C10_CONTRACT", "CHART_CONTRACT", "C10",
        len(c10_values) == 860 and c10_values["subject_key"].nunique() == 215 and observed_actual == EXPECTED_C10_OBSERVED_N and zero_actual == EXPECTED_C10_ZERO_N and c10_positive,
        {"state_rows": 860, "observed_n": EXPECTED_C10_OBSERVED_N, "zero_n": EXPECTED_C10_ZERO_N},
        {"state_rows": len(c10_values), "countries": c10_values["subject_key"].nunique(), "observed_n": observed_actual, "zero_n": zero_actual},
        "Zero remains a separate mass; the logarithmic distribution receives positive values only.",
    )

    c11 = data["entry_lag"]
    statuses = set(c11["material_entry_status_text"].dropna().astype(str))
    no_left_fake = c11.loc[c11["material_entry_status_text"].eq("MATERIAL_AT_WINDOW_START"), "observed_material_entry_year"].isna().all()
    no_right_fake = c11.loc[c11["material_entry_status_text"].eq("NOT_MATERIAL_BY_2021"), "observed_material_entry_year"].isna().all()
    collector.add(
        "P4_C11_CONTRACT", "CHART_CONTRACT", "C11",
        len(c11) == 40 and set(c11["entity"]) == set(ASEAN) and c11["series_key"].nunique() == 4 and no_left_fake and no_right_fake,
        "40 country-technology rows; censored states remain text, not boundary years",
        {"rows": len(c11), "statuses": sorted(statuses), "left_not_fabricated": no_left_fake, "right_not_fabricated": no_right_fake},
        "The renderer uses distinct symbols for left- and right-censoring.",
    )

    concentration = data["concentration"]
    growth = data["growth_contribution"]
    top3 = concentration[concentration["metric_name"].eq("top3_share")]
    top3_2021_actual = top3[top3["year"].eq(2021)].set_index("series_key")["metric_value"].to_dict()
    growth_country = growth[growth["metric_name"].eq("positive_growth_share")]
    growth_sum = float(growth_country["metric_value"].sum())
    vietnam_growth = float(growth_country.loc[growth_country["subject_key"].eq("Vietnam"), "metric_value"].iloc[0])
    top3_ok = all(round(float(top3_2021_actual[key]), 6) == round(expected, 6) for key, expected in EXPECTED_C12_2021_TOP3_SHARE.items())
    collector.add(
        "P4_C12_CONTRACT", "CHART_CONTRACT", "C12",
        len(top3) == 44 and _close(growth_sum, 1.0) and top3_ok and round(vietnam_growth, 6) == round(EXPECTED_C12_VIETNAM_POSITIVE_GROWTH_SHARE, 6),
        {"concentration_points": 44, "growth_share_sum": 1.0, "top3_2021": EXPECTED_C12_2021_TOP3_SHARE, "vietnam_growth_share": EXPECTED_C12_VIETNAM_POSITIVE_GROWTH_SHARE},
        {"concentration_points": len(top3), "growth_share_sum": growth_sum, "top3_2021": top3_2021_actual, "vietnam_growth_share": vietnam_growth},
        "Growth contribution uses gross positive additions; HHI/Gini are not substituted for it.",
    )

    c13 = data["capacity_factor"]
    row_counts = c13.groupby("technology").size().to_dict()
    countries = {technology: tuple(sorted(c13.loc[c13["technology"].eq(technology), "entity"].unique())) for technology in ["WIND", "SOLAR"]}
    recomputed = c13["generation_twh"] / (c13["capacity_gw"] * 8.76)
    max_difference = float((recomputed - c13["capacity_factor_ratio"]).abs().max())
    expected_countries = {key: tuple(sorted(value)) for key, value in CAPACITY_FACTOR_COUNTRIES.items()}
    collector.add(
        "P4_C13_CONTRACT", "CHART_CONTRACT", "C13",
        len(c13) == 104 and row_counts == EXPECTED_CAPACITY_FACTOR_ROWS and countries == expected_countries and c13["capacity_factor_ratio"].between(0, 1).all() and max_difference <= CAPACITY_FACTOR_TOLERANCE,
        {"rows": 104, "row_counts": EXPECTED_CAPACITY_FACTOR_ROWS, "countries": expected_countries, "range": [0, 1], "formula_tolerance": CAPACITY_FACTOR_TOLERANCE},
        {"rows": len(c13), "row_counts": row_counts, "countries": countries, "min": float(c13["capacity_factor_ratio"].min()), "max": float(c13["capacity_factor_ratio"].max()), "max_formula_difference": max_difference},
        "Both numerator and denominator retain source keys; year-end capacity makes this a proxy, not technical efficiency.",
    )


def phase2_fact_energy_lock_ok(input_checks: dict[str, Any]) -> bool:
    """Return whether C13's direct Phase 2 input matches every locked baseline."""
    return bool(
        input_checks["phase2_fact_rows"] == EXPECTED_PHASE2_FACT_ROWS
        and input_checks["phase2_fact_energy_sha256"] == EXPECTED_PHASE2_FACT_ENERGY_SHA256
        and input_checks["phase2_db_fingerprint"] == EXPECTED_PHASE2_FINGERPRINT
        and not input_checks["phase2_gate_error"]
    )


def enforce_phase2_fact_energy_lock(input_checks: dict[str, Any]) -> None:
    """Stop before chart-data construction when C13's direct input is stale."""
    if phase2_fact_energy_lock_ok(input_checks):
        return
    actual = {
        "rows": input_checks["phase2_fact_rows"],
        "sha256": input_checks["phase2_fact_energy_sha256"],
        "phase2_fingerprint": input_checks["phase2_db_fingerprint"],
        "gate_error": input_checks["phase2_gate_error"] or "NONE",
    }
    raise Phase4Error(f"Phase 2 fact_energy input lock failed before chart build: {actual}")


def validate_release(
    inputs: Inputs,
    input_checks: dict[str, Any],
    data: dict[str, pd.DataFrame],
    chart_dir: Path,
    workbook_path: Path,
) -> pd.DataFrame:
    """Run the locked 38 blocking checks and return their deterministic evidence."""
    collector = ValidationCollector(inputs.fingerprint)
    phase2_fact_lock_ok = phase2_fact_energy_lock_ok(input_checks)
    collector.add(
        "P4_B1_PHASE2_FACT_ENERGY", "INPUT_LOCK", "C13", phase2_fact_lock_ok,
        {"rows": EXPECTED_PHASE2_FACT_ROWS, "sha256": EXPECTED_PHASE2_FACT_ENERGY_SHA256, "phase2_fingerprint": EXPECTED_PHASE2_FINGERPRINT},
        {"rows": input_checks["phase2_fact_rows"], "sha256": input_checks["phase2_fact_energy_sha256"], "phase2_fingerprint": input_checks["phase2_db_fingerprint"], "gate_error": input_checks["phase2_gate_error"] or "NONE"},
        "C13 is bound to the exact Phase 2 fact CSV and to the independently recomputed database fingerprint.",
    )
    collector.add("P4_B1_PHASE2_VALIDATION_ROWS", "INPUT_LOCK", "NOT_APPLICABLE", input_checks["phase2_rows"] == EXPECTED_PHASE2_VALIDATION_ROWS, EXPECTED_PHASE2_VALIDATION_ROWS, input_checks["phase2_rows"], "Phase 2 validation row count is locked.")
    collector.add("P4_B1_PHASE2_STATUS_COUNTS", "INPUT_LOCK", "NOT_APPLICABLE", input_checks["phase2_status"] == EXPECTED_PHASE2_STATUS_COUNTS, EXPECTED_PHASE2_STATUS_COUNTS, input_checks["phase2_status"], "PASS/WARN/SKIP are audit statuses; only blocking failures stop release.")
    collector.add("P4_B1_PHASE2_BLOCKING", "INPUT_LOCK", "NOT_APPLICABLE", input_checks["phase2_blocking"] == EXPECTED_PHASE2_BLOCKING_FAILURES, EXPECTED_PHASE2_BLOCKING_FAILURES, input_checks["phase2_blocking"], "No blocking Phase 2 row may be non-PASS.")
    collector.add("P4_B1_PHASE3_FACT_ROWS", "INPUT_LOCK", "NOT_APPLICABLE", input_checks["phase3_fact_rows"] == EXPECTED_PHASE3_FACT_ROWS, EXPECTED_PHASE3_FACT_ROWS, input_checks["phase3_fact_rows"], "The authoritative Phase 3 fact surface is locked.")
    collector.add("P4_B1_PHASE3_QUESTION_ROWS", "INPUT_LOCK", "NOT_APPLICABLE", input_checks["phase3_question_rows"] == EXPECTED_QUESTION_ROWS, EXPECTED_QUESTION_ROWS, input_checks["phase3_question_rows"], "Q1-Q6 distribution detects a stale or partially rebuilt fact surface.")
    collector.add("P4_B1_PHASE3_CROSSCHECK", "INPUT_LOCK", "NOT_APPLICABLE", input_checks["phase3_cross_rows"] == EXPECTED_PHASE3_CROSSCHECK_ROWS and input_checks["phase3_cross_pass"] == EXPECTED_PHASE3_CROSSCHECK_ROWS, f"{EXPECTED_PHASE3_CROSSCHECK_ROWS}/{EXPECTED_PHASE3_CROSSCHECK_ROWS} PASS", f"{input_checks['phase3_cross_pass']}/{input_checks['phase3_cross_rows']} PASS", "Every Python-SQL cross-check must pass.")
    collector.add("P4_B1_PHASE3_VALIDATIONS", "INPUT_LOCK", "NOT_APPLICABLE", input_checks["phase3_validation_rows"] == EXPECTED_PHASE3_VALIDATION_ROWS and input_checks["phase3_validation_pass"] == EXPECTED_PHASE3_VALIDATION_ROWS and input_checks["workbook_contracts_ok"] and inputs.workbook_fingerprints == {EXPECTED_PHASE2_FINGERPRINT}, "92/92 PASS; six workbook contracts; one locked fingerprint", {"validation_pass": input_checks["phase3_validation_pass"], "validation_rows": input_checks["phase3_validation_rows"], "workbook_contracts": input_checks["workbook_contracts_ok"], "workbook_fingerprints": sorted(inputs.workbook_fingerprints)}, "Workbook names/sheets and provenance must agree before chart preparation.")
    collector.add("P4_B1_PHASE3_FINGERPRINT", "INPUT_LOCK", "NOT_APPLICABLE", input_checks["fingerprints"] == {EXPECTED_PHASE2_FINGERPRINT} and input_checks["archive_count"] == 17, {"fingerprint": EXPECTED_PHASE2_FINGERPRINT, "archive_csv": 17}, {"fingerprints": sorted(input_checks["fingerprints"]), "archive_csv": input_checks["archive_count"]}, "A single fingerprint binds Phase 2, Phase 3, and Phase 4.")

    chart_ids = [spec["chart_id"] for spec in CHART_SPECS]
    chart_files = [spec["file_name"] for spec in CHART_SPECS]
    collector.add("P4_B2_CHART_ALLOWLIST", "DESIGN_CONTRACT", "NOT_APPLICABLE", len(chart_ids) == 13 and len(set(chart_ids)) == 13 and len(set(chart_files)) == 13 and chart_files == list(CHART_FILES), "13 unique C01-C13 IDs and filenames", {"ids": chart_ids, "files": chart_files}, "The allow-list prevents silent chart additions, omissions, or renames.")
    metadata_ok = all(len(spec["question_ids"]) == 1 and spec["message_vi"] and spec["source_locator"] and spec["axis_mode"] and spec["limitation_vi"] for spec in CHART_SPECS)
    collector.add("P4_B2_CHART_METADATA", "DESIGN_CONTRACT", "NOT_APPLICABLE", metadata_ok, "one Q1-Q6 mapping, message, source, axis, and limitation per chart", "complete" if metadata_ok else "incomplete", "Locked metadata prevents post-hoc changes to a chart's intended claim.")
    palette_ok = set(COUNTRY_COLORS) == set(ASEAN) and set(TECHNOLOGY_COLORS) == {"HYDRO", "WIND", "SOLAR", "OTHER_RENEWABLES", "RENEWABLE_TOTAL", "MODERN_WIND_SOLAR"} and COUNTRY_COLORS["Vietnam"] == "#D55E00"
    collector.add("P4_B2_PALETTES", "DESIGN_CONTRACT", "NOT_APPLICABLE", palette_ok, "10-country and six-series coverage; Vietnam=#D55E00", {"countries": len(COUNTRY_COLORS), "series": len(TECHNOLOGY_COLORS), "vietnam": COUNTRY_COLORS.get("Vietnam")}, "Fixed colors preserve meaning across all figures.")
    forbidden = _source_has_forbidden_plot_calls()
    collector.add("P4_B2_STATIC_RENDER_RULES", "DESIGN_CONTRACT", "NOT_APPLICABLE", not forbidden and DIMENSIONS_PX == (2400, 1350) and DPI == 200 and FONT_FAMILY == "DejaVu Sans", "no pie/twinx/twiny; 2400x1350; 200 DPI; DejaVu Sans", {"forbidden_calls": forbidden, "dimensions": DIMENSIONS_PX, "dpi": DPI, "font": FONT_FAMILY}, "Static inspection enforces the non-negotiable visual grammar.")

    sheets_ok, sheet_details = _sheet_contract_result(data)
    collector.add("P4_B3_SHEET_ALLOWLIST", "TRACEABILITY", "NOT_APPLICABLE", sheets_ok, list(CHART_DATA_SHEETS), list(data), "Every plotted datum has one documented workbook home and required schema.")
    used_keys = _data_fact_keys(data)
    authoritative_keys = set(inputs.facts["fact_key"].astype(str))
    values_match, trace_evidence = _traceability_result(inputs, data)
    trace_ok = used_keys <= authoritative_keys and used_keys <= inputs.workbook_fact_keys and values_match
    trace_evidence.update({
        "missing_from_facts": len(used_keys - authoritative_keys),
        "missing_from_workbooks": len(used_keys - inputs.workbook_fact_keys),
    })
    collector.add("P4_B3_TRACEABILITY", "TRACEABILITY", "NOT_APPLICABLE", trace_ok, "all C01-C12 keys, values, units, statuses, exclusions and source versions match unique Phase 3 facts and presentation workbooks", trace_evidence, "This blocks a second, unverified version of analytical truth.")
    years_ok, actual_years = _all_years_in_domain(data)
    collector.add("P4_B3_YEAR_DOMAIN", "TRACEABILITY", "NOT_APPLICABLE", years_ok and actual_years == list(YEARS), list(YEARS), actual_years, "No chart-level observation may introduce 2022 or omit the locked window endpoints.")
    finite_ok, infinite_count = _no_unexplained_nonfinite(data)
    collector.add("P4_B3_FINITE", "TRACEABILITY", "NOT_APPLICABLE", finite_ok, 0, infinite_count, "Structural blanks retain status/exclusion fields; infinite values are never publishable.")
    unique_ok, duplicate_counts = _unique_key_result(data)
    collector.add("P4_B3_UNIQUE_KEYS", "TRACEABILITY", "NOT_APPLICABLE", unique_ok, "zero duplicate semantic keys on all 16 sheets", duplicate_counts, "Stable semantic keys prevent double plotting and ambiguous trace-back.")
    energy_pairs = set(zip(inputs.fact_energy["entity"].astype(str), pd.to_numeric(inputs.fact_energy["year"], errors="coerce"), inputs.fact_energy["metric_key"].astype(str)))
    cap = data["capacity_factor"]
    cap_source_ok = all((row.entity, float(row.year), f"{row.technology.lower()}_generation_twh") in energy_pairs and (row.entity, float(row.year), f"{row.technology.lower()}_capacity_gw") in energy_pairs for row in cap.itertuples())
    collector.add("P4_B3_CAPACITY_SOURCE_KEYS", "TRACEABILITY", "C13", cap_source_ok and cap["generation_source_key"].notna().all() and cap["capacity_source_key"].notna().all(), "two resolvable source keys for every C13 row", {"rows": len(cap), "generation_keys": int(cap["generation_source_key"].notna().sum()), "capacity_keys": int(cap["capacity_source_key"].notna().sum())}, "C13 is the only Phase 4 calculation and therefore retains both operands' source keys.")

    _validate_chart_contracts(collector, inputs, data)

    actual_pngs = sorted(path.name for path in chart_dir.glob("*.png"))
    collector.add("P4_B9_PNG_ALLOWLIST", "RELEASE", "NOT_APPLICABLE", set(actual_pngs) == set(CHART_FILES) and len(actual_pngs) == 13, sorted(CHART_FILES), actual_pngs, "Unexpected PNG files are a blocking error and are never deleted automatically.")
    dimensions: dict[str, list[int]] = {}
    nonblank: dict[str, float] = {}
    image_open_ok = True
    for file_name in CHART_FILES:
        path = chart_dir / file_name
        try:
            with Image.open(path) as image:
                image.load()
                dimensions[file_name] = [image.width, image.height]
                stddev = ImageStat.Stat(image.convert("RGB")).stddev
                nonblank[file_name] = float(max(stddev))
                image_open_ok = image_open_ok and image.format == "PNG"
        except Exception:
            image_open_ok = False
            dimensions[file_name] = [0, 0]
            nonblank[file_name] = 0.0
    dimensions_ok = image_open_ok and all(tuple(value) == DIMENSIONS_PX for value in dimensions.values())
    collector.add("P4_B9_PNG_DIMENSIONS", "RELEASE", "NOT_APPLICABLE", dimensions_ok, {name: list(DIMENSIONS_PX) for name in CHART_FILES}, dimensions, "Decoded dimensions, not file metadata, enforce the 2400x1350 contract.")
    nonblank_ok = image_open_ok and all(value > 0.5 for value in nonblank.values())
    collector.add("P4_B9_PNG_NONBLANK", "RELEASE", "NOT_APPLICABLE", nonblank_ok, "all decoded RGB standard deviations > 0.5", nonblank, "Openable, non-white images are required before publication.")
    expected_manifest = pd.DataFrame(manifest_rows(), columns=MANIFEST_COLUMNS)
    manifest_ok = data["manifest"].astype(str).equals(expected_manifest.astype(str)) and len(data["manifest"]) == 13
    collector.add("P4_B9_MANIFEST", "RELEASE", "NOT_APPLICABLE", manifest_ok, "13 exact locked manifest rows", {"rows": len(data["manifest"]), "columns": list(data["manifest"].columns)}, "Manifest equality binds filenames, questions, messages, sources, limits, and fingerprint.")
    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        workbook_schema_ok = tuple(workbook.sheetnames) == tuple(CHART_DATA_SHEETS)
        workbook_schema_ok = workbook_schema_ok and set(data["metadata"]["metadata_key"]) == set(REQUIRED_METADATA_KEYS)
        workbook_schema_ok = workbook_schema_ok and set(data["method_notes"]["note_key"]) == set(REQUIRED_METHOD_NOTE_KEYS)
        workbook_sheets = list(workbook.sheetnames)
    finally:
        workbook.close()
    collector.add("P4_B9_WORKBOOK_SCHEMA", "RELEASE", "NOT_APPLICABLE", workbook_schema_ok, list(CHART_DATA_SHEETS), workbook_sheets, "The evidence workbook must be complete and ordered for deterministic review.")
    current_hashes = file_hashes(Path(path) for path in inputs.immutable_hashes)
    immutable_ok = current_hashes == inputs.immutable_hashes
    changed = sorted(path for path in set(current_hashes) | set(inputs.immutable_hashes) if current_hashes.get(path) != inputs.immutable_hashes.get(path))
    collector.add("P4_B9_SOURCE_IMMUTABILITY", "RELEASE", "NOT_APPLICABLE", immutable_ok, "17 archive CSV, fact_energy.csv, and all protected Phase 3 inputs unchanged", changed or "unchanged", "Hash comparison prevents chart generation from mutating its authoritative inputs.")

    result = collector.frame()
    if len(result) != len(EXPECTED_VALIDATION_IDS) or set(result["validation_id"]) != set(EXPECTED_VALIDATION_IDS):
        raise Phase4Error("Phase 4 validation ID contract is incomplete or duplicated")
    return result


def write_chart_workbook(data: dict[str, pd.DataFrame], path: Path) -> None:
    """Write deterministic evidence sheets with review-friendly formatting."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name in CHART_DATA_SHEETS:
            data[sheet_name].to_excel(writer, sheet_name=sheet_name, index=False)
    workbook = load_workbook(path)
    try:
        header_fill = PatternFill("solid", fgColor="1F4E78")
        for worksheet in workbook.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for cell in worksheet[1]:
                cell.font = Font(color="FFFFFF", bold=True)
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for index, column_cells in enumerate(worksheet.iter_cols(min_row=1, max_row=min(worksheet.max_row, 101)), start=1):
                width = min(60, max(10, max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells) + 2))
                worksheet.column_dimensions[column_cells[0].column_letter].width = width
        workbook.save(path)
    finally:
        workbook.close()


def write_validation_csv(frame: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    frame.to_csv(path, index=False, encoding="utf-8-sig", lineterminator="\n")


def replace_outputs_atomically(staged_targets: Iterable[tuple[Path, Path]]) -> None:
    """Publish every staged file or restore the exact prior target set on failure."""
    pairs = [(Path(staged), Path(target)) for staged, target in staged_targets]
    if not pairs:
        return
    for staged, target in pairs:
        if not staged.is_file():
            raise Phase4Error(f"Missing staged publication member: {staged}")
        target.parent.mkdir(parents=True, exist_ok=True)
    backup_root = Path(tempfile.mkdtemp(prefix=".phase4_backup_", dir=str(pairs[0][1].parent)))
    remove_backup_root = True
    existed: dict[Path, bool] = {}
    backups: dict[Path, Path] = {}
    published: list[Path] = []
    try:
        for index, (_, target) in enumerate(pairs):
            existed[target] = target.exists()
            if existed[target]:
                backup = backup_root / f"{index:03d}_{target.name}.bak"
                shutil.copy2(target, backup)
                backups[target] = backup
        for staged, target in pairs:
            os.replace(staged, target)
            published.append(target)
    except Exception as exc:
        rollback_errors: list[str] = []
        for target in reversed(published):
            try:
                if existed.get(target, False):
                    backup = backups[target]
                    if backup.exists():
                        os.replace(backup, target)
                elif target in published and target.exists():
                    target.unlink()
            except Exception as rollback_exc:  # pragma: no cover - catastrophic disk failure
                rollback_errors.append(f"{target}: {rollback_exc}")
        if rollback_errors:
            remove_backup_root = False
            message = (
                "Phase 4 publication failed; rollback incomplete; "
                f"backup retained at {backup_root.resolve()}; publication error: {exc}; "
                f"rollback errors={rollback_errors}"
            )
        else:
            message = f"Phase 4 publication failed; previous outputs restored: {exc}"
        raise Phase4Error(message) from exc
    finally:
        if remove_backup_root:
            shutil.rmtree(backup_root, ignore_errors=True)


def _official_unknown_pngs() -> list[str]:
    if not CHARTS.exists():
        return []
    return sorted(path.name for path in CHARTS.glob("*.png") if path.name not in set(CHART_FILES))


def build_stage() -> tuple[Path, list[tuple[Path, Path]], pd.DataFrame]:
    """Build and fully validate a temporary release without touching official outputs."""
    unexpected = _official_unknown_pngs()
    if unexpected:
        raise Phase4Error(f"Unexpected PNG files in output/charts; review manually, none were deleted: {unexpected}")
    inputs, input_checks = load_inputs()
    enforce_phase2_fact_energy_lock(input_checks)
    data = build_chart_data(inputs)
    stage_root = Path(tempfile.mkdtemp(prefix=".phase4_stage_", dir=str(OUTPUT)))
    stage_charts = stage_root / "charts"
    stage_stats = stage_root / "stats"
    workbook_path = stage_stats / "04_chart_data.xlsx"
    validation_path = stage_stats / "04_validation_results.csv"
    try:
        stage_charts.mkdir(parents=True, exist_ok=True)
        stage_stats.mkdir(parents=True, exist_ok=True)
        write_chart_workbook(data, workbook_path)
        render_charts(data, stage_charts)
        validation = validate_release(inputs, input_checks, data, stage_charts, workbook_path)
        write_validation_csv(validation, validation_path)
        failures = validation[validation["status"].ne("PASS")]
        if not failures.empty:
            failed_ids = failures["validation_id"].tolist()
            raise Phase4Error(f"Blocking Phase 4 validations failed: {failed_ids}; staged evidence retained only until cleanup")
        pairs = [
            *((stage_charts / file_name, CHARTS / file_name) for file_name in CHART_FILES),
            (workbook_path, STATS / "04_chart_data.xlsx"),
            (validation_path, STATS / "04_validation_results.csv"),
        ]
        return stage_root, pairs, validation
    except Exception:
        shutil.rmtree(stage_root, ignore_errors=True)
        raise


def main() -> int:
    stage_root: Path | None = None
    try:
        stage_root, pairs, validation = build_stage()
        replace_outputs_atomically(pairs)
        print(
            "PHASE 4 CHART RELEASE: PASS | "
            f"charts={len(CHART_FILES)} | workbook_sheets={len(CHART_DATA_SHEETS)} | "
            f"validations={len(validation)} | fingerprint={EXPECTED_PHASE2_FINGERPRINT}"
        )
        return 0
    except Exception as exc:
        print(f"PHASE 4 CHART RELEASE: FAIL | {exc}", file=sys.stderr)
        return 1
    finally:
        if stage_root is not None:
            shutil.rmtree(stage_root, ignore_errors=True)


if __name__ == "__main__":
    raise SystemExit(main())
