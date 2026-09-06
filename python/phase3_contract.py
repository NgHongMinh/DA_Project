"""Locked release contract shared by Phase 3 build and verification."""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


EXPECTED_FACT_ROWS = 18_926
EXPECTED_QUESTION_ROWS = {
    "Q1": 5_544,
    "Q2": 3_608,
    "Q3": 1_980,
    "Q4": 1_382,
    "Q5": 5_950,
    "Q6": 462,
}
EXPECTED_CROSSCHECK_ROWS = 18_926
EXPECTED_PHASE3_VALIDATION_ROWS = 92
RATIO_NUMBER_FORMAT = "0.0000%"

PHASE3_WORKBOOKS = {
    "03_global_cross_section.xlsx": [
        "generation_stats", "country_values", "coverage", "metadata",
    ],
    "03_asean_cross_section.xlsx": [
        "generation_stats", "asean_annual", "country_year_values",
        "twh_vs_pct", "endpoint_change", "metadata",
    ],
    "03_vietnam_time_series.xlsx": [
        "annual", "annual_changes", "asean9_benchmark", "endpoint_change",
        "metadata",
    ],
    "03_ranking_comparison.xlsx": [
        "annual_ranks", "milestone_change", "cagr_ranking", "metadata",
    ],
    "03_entry_lag.xlsx": [
        "entry_status", "entry_summary", "method_notes", "metadata",
    ],
    "03_global_context.xlsx": [
        "annual_context", "period_growth", "annual_concentration",
        "growth_contribution", "world_reconciliation", "metadata",
    ],
}

_BASE_VALIDATION_IDS = {
    "P3_PHASE2_VALIDATION_GATE", "P3_INPUT", "P3_KEYS", "P3_FINITE",
    "P3_CODES", "P3_STATUS", "P3_CROSS", "P3_DEFINITIONS",
    "P3_Q5_PAIR_MAPPING", "P3_DEFINITION_UNITS", "P3_FINGERPRINT",
    "P3_PERIOD_KEYS", "P3_COMPONENT_COUNTS", "P3_ASEAN_PANEL",
    "P3_TECH_ROWS", "P3_TECH_SHARE", "P3_COUNTRY_SHARE",
    "P3_PERCENT_DOMAINS", "P3_Q5_COVERAGE", "P3_YOY_RULE",
    "P3_CAGR_RULE", "P3_RANK_COVERAGE", "P3_ANNUAL_RANK_RULE",
    "P3_RANK_CHANGE_RULE", "P3_CONCENTRATION_DOMAIN", "P3_CONTRIBUTION",
    "P3_CONTRIBUTION_COHORT", "P3_DELTA_RECONCILIATION",
    "P3_RELATIVE_LAG", "P3_CANONICAL_CONCENTRATION",
    "P3_QUESTION_COVERAGE", "P3_VIETNAM_ANNUAL",
    "P3_DISTRIBUTION_WORKBOOK_CONTRACT", "P3_Q6_WORKBOOK_SUMMARY",
    "P3_WORKBOOK_PROVENANCE", "P3_VALIDATION_IDS_UNIQUE",
}


def _workbook_validation_ids() -> set[str]:
    identifiers: set[str] = set()
    for workbook_name, sheet_names in PHASE3_WORKBOOKS.items():
        workbook_id = Path(workbook_name).stem.upper()
        identifiers.add(f"WB_{workbook_id}")
        identifiers.add(f"META_{workbook_id}")
        for sheet_name in sheet_names:
            if sheet_name in {"metadata", "method_notes"}:
                continue
            sheet_id = sheet_name.upper()
            identifiers.add(f"TRACE_{workbook_id}_{sheet_id}")
            identifiers.add(f"KEY_{workbook_id}_{sheet_id}")
    return identifiers


EXPECTED_PHASE3_VALIDATION_IDS = frozenset(
    _BASE_VALIDATION_IDS | _workbook_validation_ids()
)
if len(EXPECTED_PHASE3_VALIDATION_IDS) != EXPECTED_PHASE3_VALIDATION_ROWS:
    raise RuntimeError("Phase 3 validation ID contract must contain exactly 92 IDs")


def workbook_number_format_errors(path: Path, sheet_names: list[str]) -> list[str]:
    """Return ratio/pct number-format contract violations in one workbook."""
    errors: list[str] = []
    workbook = load_workbook(path, data_only=False, read_only=False)
    try:
        if workbook.sheetnames != sheet_names:
            errors.append(
                f"sheet order expected={sheet_names!r} actual={workbook.sheetnames!r}"
            )
            return errors
        for sheet_name in sheet_names:
            worksheet = workbook[sheet_name]
            headers = {
                cell.value: cell.column
                for cell in worksheet[1]
                if isinstance(cell.value, str)
            }
            unit_pairs: list[tuple[str, str]] = []
            if "unit" in headers:
                unit_pairs.append(("metric_value", "unit"))
            unit_pairs.extend(
                (unit_name[:-5], unit_name)
                for unit_name in headers
                if unit_name.endswith("_unit") and unit_name[:-5] in headers
            )
            for value_name, unit_name in unit_pairs:
                value_column = headers[value_name]
                unit_column = headers[unit_name]
                for row_number in range(2, worksheet.max_row + 1):
                    unit = worksheet.cell(row=row_number, column=unit_column).value
                    value_cell = worksheet.cell(row=row_number, column=value_column)
                    if value_cell.value is None:
                        continue
                    if unit == "ratio" and value_cell.number_format != RATIO_NUMBER_FORMAT:
                        errors.append(
                            f"{sheet_name}!{value_cell.coordinate}: ratio format "
                            f"{value_cell.number_format!r}"
                        )
                    elif unit == "pct" and "%" in value_cell.number_format:
                        errors.append(
                            f"{sheet_name}!{value_cell.coordinate}: pct must remain on the 0-100 scale"
                        )
    finally:
        workbook.close()
    return errors
