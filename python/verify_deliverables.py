"""Đối chiếu dữ liệu và nội dung bàn giao; dựng bản xem trước bố cục từ tệp đã lưu.

Ảnh xem trước dùng Pillow, không phải kết xuất bằng PowerPoint.
"""
from __future__ import annotations
import contextlib
import hashlib
import io
import json
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.text import Text
import numpy as np
import pandas as pd
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.text import PP_ALIGN

from config import ARCHIVE, CLEAN, OUTPUT
from publication_checks import assert_notebook_clean, assert_reader_text, assert_slides_clean


def preview(prs, directory):
    directory.mkdir(parents=True, exist_ok=True)
    scale = 1600 / prs.slide_width
    pixels_per_point = scale * 12700
    thumbnails = []
    for page, slide in enumerate(prs.slides, 1):
        canvas = Image.new("RGB", (1600, 900), "white")
        draw = ImageDraw.Draw(canvas)
        for shape in slide.shapes:
            x, y, w, h = [round(v*scale) for v in (shape.left, shape.top, shape.width, shape.height)]
            if shape.shape_type == 13:
                image = Image.open(io.BytesIO(shape.image.blob)).convert("RGB").resize((w, h), Image.Resampling.LANCZOS)
                canvas.paste(image, (x, y))
                continue
            if shape.shape_type == 1 and shape.fill.type == 1:
                draw.rectangle((x, y, x+w, y+h), fill="#"+str(shape.fill.fore_color.rgb))
            if not shape.has_text_frame or not shape.text.strip():
                continue
            frame = shape.text_frame
            top = y + frame.margin_top*scale
            for paragraph in frame.paragraphs:
                if not paragraph.runs:
                    continue
                run = paragraph.runs[0]
                size = run.font.size.pt if run.font.size else 18
                font_path = Path("C:/Windows/Fonts") / ("segoeuib.ttf" if run.font.bold else "segoeui.ttf")
                font = ImageFont.truetype(str(font_path), max(1, round(size*pixels_per_point)))
                length = draw.textlength(paragraph.text, font=font)
                left = x + frame.margin_left*scale
                if paragraph.alignment == PP_ALIGN.CENTER:
                    left = x+(w-length)/2
                elif paragraph.alignment == PP_ALIGN.RIGHT:
                    left = x+w-frame.margin_right*scale-length
                draw.text((left, top), paragraph.text, font=font, fill="#"+str(run.font.color.rgb), anchor="lt")
                spacing = paragraph.line_spacing
                top += spacing.pt*pixels_per_point if hasattr(spacing, "pt") else size*pixels_per_point*1.24
            if top > y+h+2:
                raise AssertionError(f"Trang {page}: chữ vượt chiều cao khung")
        canvas.save(directory / f"page-{page:02d}.png")
        thumb = canvas.resize((480, 270), Image.Resampling.LANCZOS)
        thumbnails.append(thumb)
    contact = Image.new("RGB", (1440, 6*296), "#E5E7EB")
    draw = ImageDraw.Draw(contact)
    font = ImageFont.truetype("C:/Windows/Fonts/segoeui.ttf", 17)
    for i, thumb in enumerate(thumbnails):
        x, y = (i % 3)*480, (i // 3)*296
        draw.text((x+8, y+3), f"Trang {i+1}", fill="#374151", font=font)
        contact.paste(thumb, (x, y+26))
    contact.save(directory / "overview.png")


def main():
    path = OUTPUT / "notebook" / "ASEAN_Renewables_DataProcessing.ipynb"
    notebook = json.loads(path.read_text(encoding="utf-8"))
    assert_notebook_clean(notebook)
    code = [c for c in notebook["cells"] if c["cell_type"] == "code"]
    assert len(code) == 31
    assert [c["execution_count"] for c in code] == list(range(1, 32))
    assert not any(o["output_type"] == "error" for c in code for o in c.get("outputs", []))
    assert sum("image/png" in o.get("data", {}) for c in code for o in c.get("outputs", [])) == 5
    hashes = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in ARCHIVE.glob("*.csv")}
    figures = []

    def capture():
        for number in plt.get_fignums():
            fig = plt.figure(number)
            fig.canvas.draw()
            labels = [t.get_text() for t in fig.findobj(Text) if t.get_text()]
            figures.append(labels)
            assert_reader_text([("Biểu đồ trong sổ tay", "\n".join(labels))])
        plt.close("all")

    namespace = {"__name__": "__main__"}
    original_show = plt.show
    plt.show = capture
    try:
        # Chạy đúng mã lưu trong tệp, không đọc kết quả tính sẵn vào phép biến đổi.
        with contextlib.redirect_stdout(io.StringIO()):
            for cell in code:
                source = cell["source"]
                if not isinstance(source, str):
                    source = "".join(source)
                exec(compile(source, "<ô sổ tay>", "exec"), namespace)
    finally:
        plt.show = original_show
    assert len(figures) == 5
    actual = namespace["fact_energy"].sort_values(["entity", "year", "metric_key"]).reset_index(drop=True)
    expected = pd.read_csv(CLEAN / "fact_energy.csv").sort_values(["entity", "year", "metric_key"]).reset_index(drop=True)
    assert len(actual) == len(expected) == 89075
    pd.testing.assert_frame_equal(actual, expected[actual.columns], check_dtype=False, rtol=0, atol=1e-9)
    difference = float(np.max(np.abs(actual.value-expected.value)))
    at_threshold = actual.loc[actual.metric_key.str.endswith("_twh") & actual.value.eq(.05)]
    assert len(at_threshold) and at_threshold.value_class.eq("Material").all()
    zero_rate, reason = namespace["safe_growth_rate"](10, 0)
    assert pd.isna(zero_rate) and reason != "NONE"
    namespace["conn"].close()
    assert hashes == {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in ARCHIVE.glob("*.csv")}
    slides = OUTPUT / "slides"
    prs = Presentation(slides / "ASEAN_Renewables.pptx")
    assert len(prs.slides) == 17 and abs(prs.slide_width/prs.slide_height-16/9) < .001
    labels = json.loads((slides / "assets" / "figure_text.json").read_text(encoding="utf-8"))
    assert_slides_clean(prs, labels)
    core = prs.core_properties
    assert not core.comments and core.last_modified_by == core.author == "Nguyễn Hồng Minh"
    for slide in prs.slides:
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left+shape.width <= prs.slide_width+2 and shape.top+shape.height <= prs.slide_height+2
            if shape.shape_type == 13:
                with Image.open(io.BytesIO(shape.image.blob)) as image:
                    assert_reader_text([("Thuộc tính ảnh", str(image.info))])
    preview(prs, slides / "previews")
    evidence = pd.read_csv(OUTPUT / "stats" / "06_finding_evidence.csv")
    assert len(evidence) == 77 and evidence.status.eq("MATCH").all()
    report = {
        "slides": 17, "code_cells_executed": 31, "notebook_figures": 5,
        "matched_claims": 77, "fact_rows": len(actual), "max_absolute_difference": difference,
        "source_files_unchanged": len(hashes), "threshold_rows_checked": len(at_threshold),
        "internal_content_check": "passed", "preview_engine": "Pillow; not PowerPoint rendering",
    }
    (slides / "previews" / "verification.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
